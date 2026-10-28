"""Software contracts only; no patient cohort or real Notes workbook."""
import unittest
import csv
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from serda_pipeline.ga_extraction import extract_text, reconstruct_fragments, resolve_evidence, run
from serda_pipeline.notes import schema_mapping
from serda_pipeline.core import PipelineError


class GAExtractionContracts(unittest.TestCase):
    def test_header_only_independent_outputs(self):
        with TemporaryDirectory() as temp:
            root=Path(temp)
            source=root/'source';source.mkdir()
            output=root/'output';output.mkdir()
            for name,headers in [('Frailty Demographics.xlsx',['PAT_ID','CONTACT_DATE']),
                                 ('Notes.xlsx',['PAT_ID','NOTE_ID','LINE','NOTE_CONTACT_DATE','NOTE_TEXT','Anchor','NOTE_TYPE'])]:
                book=Workbook();book.active.append(headers);book.save(source/name)
            config={'files':{},'sheets':{},'header_rows':{},'columns':{},
                    'excel_date_min':20000,'excel_date_max':80000,'dayfirst':False,
                    'notes':{'note_key_fields':['note_id'],'note_key_verified':True}}
            result=run(source,output,config)
            self.assertEqual(result['reconstructed_notes'],0)
            for name in ['ga_note_selection_audit.csv','ga_extracted_scores.csv',
                         'ga_extraction_review_queue.csv','ga_domain_coverage.csv',
                         'ga_extraction_summary.md','ga_schema_validation.json']:
                self.assertTrue((output/name).is_file(),name)

    def test_unverified_note_identity_blocks_extraction(self):
        with TemporaryDirectory() as temp:
            root=Path(temp);source=root/'source';source.mkdir();output=root/'output';output.mkdir()
            demographics=Workbook();demographics.active.append(['PAT_ID','CONTACT_DATE'])
            demographics.active.append(['TEST_ONLY','2024-01-02'])
            demographics.save(source/'Frailty Demographics.xlsx')
            notes=Workbook()
            notes.active.append(['PAT_ID','NOTE_CSN','LINE','NOTE_CONTACT_DATE','NOTE_TEXT','Anchor','NOTE_TYPE'])
            notes.active.append(['TEST_ONLY','CSN_TEST',1,'2024-01-02',
                                 'Geriatric components MoCA 15/30','Yes','Geriatric assessment'])
            notes.save(source/'Notes.xlsx')
            config={'files':{},'sheets':{},'header_rows':{},'columns':{},
                    'excel_date_min':20000,'excel_date_max':80000,'dayfirst':False,
                    'notes':{'note_key_fields':['note_csn'],'note_key_verified':False}}
            result=run(source,output,config)
            self.assertFalse(result['identity_verified'])
            self.assertEqual(result['excluded_identity_unverified'],1)
            with (output/'ga_extracted_scores.csv').open(newline='',encoding='utf-8') as handle:
                self.assertEqual(len(list(csv.DictReader(handle))),0)

    def test_single_test_fixture_selects_note_without_predictor_output(self):
        # Software fixture only: no clinical validation or synthetic cohort result.
        with TemporaryDirectory() as temp:
            root=Path(temp);source=root/'source';source.mkdir();output=root/'output';output.mkdir()
            demographics=Workbook()
            demographics.active.append(['PAT_ID','PAT_MRN_ID','CONTACT_DATE'])
            demographics.active.append(['TEST_ONLY','TEST_MRN','2024-01-02'])
            demographics.save(source/'Frailty Demographics.xlsx')
            notes=Workbook()
            notes.active.append(['PAT_ID','PAT_MRN_ID','NOTE_ID','LINE','NOTE_CONTACT_DATE','NOTE_TEXT','Anchor','NOTE_TYPE'])
            notes.active.append(['TEST_ONLY','TEST_MRN','NOTE_TEST',1,'2024-01-02','Geriatric components','Yes','Geriatric assessment'])
            notes.active.append(['TEST_ONLY','TEST_MRN','NOTE_TEST',2,'2024-01-02','MoCA 15/30','Yes','Geriatric assessment'])
            notes.save(source/'Notes.xlsx')
            config={'files':{},'sheets':{},'header_rows':{},'columns':{},
                    'excel_date_min':20000,'excel_date_max':80000,'dayfirst':False,
                    'notes':{'note_key_fields':['note_id'],'note_key_verified':True}}
            result=run(source,output,config,'mrn')
            self.assertEqual(result['confirmed_by_configured_rules'],1)
            with (output/'ga_extracted_scores.csv').open(newline='',encoding='utf-8') as handle:
                records=list(csv.DictReader(handle))
            self.assertEqual([(r['domain_or_instrument'],r['value'],r['extraction_status']) for r in records],
                             [('moca','15','extracted')])
            self.assertFalse((output/'patient_features.csv').exists())
            notes.active.append(['TEST_ONLY','TEST_MRN','NOTE_TEST_2',1,'2024-01-02',
                                 'Geriatric components','Yes','Geriatric assessment'])
            notes.active.append(['TEST_ONLY','TEST_MRN','NOTE_TEST_2',2,'2024-01-02',
                                 'GDS 2/15','Yes','Geriatric assessment'])
            notes.save(source/'Notes.xlsx')
            ambiguous_output=root/'ambiguous_output';ambiguous_output.mkdir()
            ambiguous=run(source,ambiguous_output,config,'mrn')
            self.assertEqual(ambiguous['ambiguous_multiple_ga_notes'],2)
            self.assertEqual(ambiguous['confirmed_by_configured_rules'],0)

    def test_reconstruct_line_order_and_duplicate(self):
        rows = [(0,1,'1','MoCA '),(1,2,'2','15/30'),
                (2,1,'1','MoCA '),(3,2,'2','15/30')]
        text, reasons, quality = reconstruct_fragments(rows,1,'')
        self.assertEqual(text,'MoCA 15/30')
        self.assertEqual(reasons,[])
        self.assertEqual(quality['duplicate_line_rows'],2)

    def test_reject_conflict_and_invalid_line(self):
        for rows, reason in [
            ([(0,1,'1','a'),(1,1,'1','b')],'conflicting_line_text'),
            ([(0,None,'x','a')],'missing_or_nonnumeric_line'),
        ]:
            with self.subTest(reason=reason):
                text, reasons, _ = reconstruct_fragments(rows)
                self.assertIsNone(text)
                self.assertIn(reason,reasons)

    def test_sortable_order_is_recovered_and_gap_is_review_evidence(self):
        text,reasons,quality=reconstruct_fragments([(0,3,'3','c'),(1,1,'1','a'),(2,2,'2','b')])
        self.assertEqual(text,'a\nb\nc')
        self.assertEqual(reasons,[])
        self.assertIn('source_order_sorted_by_line',quality['review_warnings'])
        text,reasons,quality=reconstruct_fragments([(0,1,'1','a'),(1,3,'3','c')])
        self.assertEqual(text,'a\nc')
        self.assertEqual(reasons,[])
        self.assertIn('line_gap_or_start',quality['review_warnings'])

    def test_explicit_numeric_instruments(self):
        rows = resolve_evidence(extract_text('MoCA 15/30. GDS 2/15. SPPB: 9/12. VES-13 score 4.')[0])
        self.assertEqual([(r[0],r[1],r[2],r[3]) for r in rows],
                         [('moca',15,30,'extracted'),('gds',2,15,'extracted'),
                          ('sppb',9,12,'extracted'),('ves_13',4,None,'extracted')])

    def test_history_negation_and_conflict(self):
        historical = resolve_evidence(extract_text('Previous MoCA 10/30; today MoCA 15/30.')[0])
        self.assertEqual([r[3] for r in historical],['historical_context','extracted'])
        self.assertEqual(resolve_evidence(extract_text('MoCA 10/30. MoCA 15/30.')[0])[0][3],
                         'ambiguous_conflicting_values')
        self.assertEqual(resolve_evidence(extract_text('MoCA 40/30.')[0])[0][3],
                         'invalid_numeric_bound')
        self.assertEqual(resolve_evidence(extract_text('MoCA 15/20.',{'moca':[30]})[0])[0][3],
                         'unverified_denominator')
        self.assertEqual(resolve_evidence(extract_text('MoCA not performed.')[0])[0][3],
                         'unsupported_format')
        self.assertEqual(resolve_evidence(extract_text('No prior falls.')[0])[0][1],
                         'no_falls_documented')
        self.assertEqual(resolve_evidence(extract_text('Denies history of falls.')[0])[0][3],
                         'negated_or_uncertain_context')

    def test_missing_schema_and_ambiguous_identity_mapping(self):
        headers,mapping,ambiguous,absent = schema_mapping(['NOTE_CSN','LINE','NOTE_TEXT'],{})
        self.assertIn('date',absent)
        self.assertIn('anchor_flag',absent)
        self.assertEqual(mapping['note_csn'],0)
        with self.assertRaises(PipelineError):
            schema_mapping(['NOTE_CSN','note csn','LINE','NOTE_TEXT'],{})


if __name__ == '__main__': unittest.main()
