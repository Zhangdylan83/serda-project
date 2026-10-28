"""MRN linkage audit for all eight exports; no cohort or feature mutation.

Rows are streamed into SQLite under the approved SERDA output directory. The
stage contains identifiers but no note text, and is removed after the run.
"""
from collections import Counter, defaultdict
from contextlib import contextmanager, closing
import csv
from datetime import datetime
import json
from pathlib import Path
import re
import sqlite3
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
import pandas as pd
import xlrd

from .core import PipelineError, parse_dates
from .notes import anchor_status, identifier, missing_value, line_number
from .schema import STEMS, SCHEMAS, normalize

DATASETS = [*STEMS, 'notes']
MRN_ALIASES = ['PAT_MRN_ID', 'MRN', 'PAT_MRN', 'PATIENT_MRN', 'PATIENT_MRN_ID',
               'MRN_ID', 'MEDICAL_RECORD_NUMBER', 'MEDICAL_RECORD_NUM']
OUTPUTS = {
 'mrn_schema_validation.csv': ['dataset','file','sheet','header_row','source_columns','pat_id_column',
      'mrn_column','other_mrn_columns','date_column','mrn_mapping_status','row_count',
      'missing_mrn_rows','numeric_mrn_cell_rows','numeric_alternative_mrn_rows','conflicting_mrn_columns_rows',
      'alternate_present_preferred_missing_rows','trimmed_mrn_rows','invalid_date_rows','reason'],
 'mrn_patient_mapping_audit.csv': ['dataset','rows','unique_pat_ids','unique_mrns','missing_mrn_rows',
      'pat_ids_with_missing_mrn','pat_ids_all_rows_missing_mrn','rows_missing_both_ids',
      'mrns_with_multiple_pat_ids','pat_ids_with_multiple_mrns',
      'mrn_values_with_case_variants','numeric_mrn_cell_rows','conflicting_mrn_columns_rows','status'],
 'mrn_dataset_inventory.csv': ['dataset','total_rows','exact_duplicate_rows','unique_pat_ids',
      'unique_mrns','missing_mrn_rows','pat_ids_with_missing_mrn','pat_ids_all_rows_missing_mrn','status'],
 'mrn_patient_intersections.csv': ['dataset','pat_id_patients','mrn_patients','pat_id_overlap_demographics',
      'mrn_overlap_demographics','mrn_outside_demographics','demographics_mrns_without_records',
      'demographics_mrn_coverage_fraction','status'],
 'mrn_coverage_comparison.csv': ['scope','dataset_a','dataset_b','pat_id_count','mrn_count',
      'pat_id_overlap','mrn_overlap','pat_id_overlap_fraction','mrn_overlap_fraction','status'],
 'mrn_anchor_alignment_summary.csv': ['metric','pat_id_count','mrn_count','status'],
 'mrn_multiple_anchor_summary.csv': ['metric','pat_id_count','mrn_count','status'],
 'mrn_vs_pat_id_comparison.csv': ['dataset','total_rows','pat_id_patients','mrn_patients',
      'patient_count_change_mrn_minus_pat_id','pat_id_demographics_overlap',
      'mrn_demographics_overlap','pat_id_coverage_fraction','mrn_coverage_fraction','status'],
 'mrn_temporal_summary.csv': ['dataset','matched_mrn_rows','ambiguous_reference_anchor_rows',
      'dated_records','missing_dates','post_anchor_records','lookback_records','anchor_day_excluded_rows','status'],
 'mrn_pairwise_intersections.csv': ['dataset_a','dataset_b','pat_id_overlap','mrn_overlap',
      'pat_id_jaccard','mrn_jaccard','potential_case_only_mrn_matches','status'],
 'mrn_pat_ids_per_mrn_distribution.csv': ['dataset','PAT_IDs_per_MRN','MRNs'],
 'mrn_mrns_per_pat_id_distribution.csv': ['dataset','MRNs_per_PAT_ID','PAT_IDs'],
 'mrn_records_per_patient_summary.csv': ['dataset','patient_key','linked_patients',
      'patients_with_multiple_rows','median_rows','max_rows','status'],
}


def write_csv(path, columns, records):
    with path.open('w',newline='',encoding='utf-8') as handle:
        writer=csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(records)


def _clean_id(value):
    value=missing_value(value)
    return identifier(value) if value is not None else None


def _clean_mrn(value):
    if value is None or (isinstance(value,str) and (not value.strip() or value in {'NULL','null','(null)','None'})):
        return None, 'missing', False
    if not isinstance(value,str):
        # An Excel numeric cell may already have lost significant leading zeroes.
        return None, 'numeric_or_nontext_cell', False
    cleaned=value.strip()
    return cleaned, 'available', cleaned != value


def _field_index(headers, choices):
    matches=[i for i,h in enumerate(headers) if normalize(h) in {normalize(a) for a in choices}]
    return matches[0] if len(matches)==1 else None


def _discover(headers, dataset, config):
    if not headers or any(h is None or not str(h).strip() for h in headers):
        raise PipelineError('Blank header in '+dataset+'; configure the exact header row.')
    names=[str(h) for h in headers]
    if len(set(map(normalize,names))) != len(names):
        raise PipelineError('Duplicate or normalization-colliding headers in '+dataset+'.')
    overrides=config.get('columns',{}).get(dataset,{})
    patient=_field_index(names,[overrides['patient']] if 'patient' in overrides else ['PAT_ID'])
    all_mrn=[i for i,h in enumerate(names) if normalize(h) in {normalize(a) for a in MRN_ALIASES}]
    if 'mrn' in overrides:
        primary=_field_index(names,[overrides['mrn']])
        status='configured' if primary is not None else 'configured_mrn_missing_or_ambiguous'
    else:
        preferred=_field_index(names,['PAT_MRN_ID'])
        if preferred is not None:
            primary=preferred; status='preferred_pat_mrn_id'
        elif len(all_mrn)==1:
            primary=all_mrn[0]; status='single_known_alias'
        elif not all_mrn:
            primary=None; status='mrn_column_missing'
        else:
            primary=None; status='multiple_mrn_columns_configure_override'
    if primary==patient and primary is not None:
        primary=None; status='mrn_column_is_pat_id_column'
    alternatives=[i for i in all_mrn if i!=primary and i!=patient]
    schema=SCHEMAS.get(dataset,{})
    date_field='anchor' if dataset=='demographics' else 'date'
    date_alias=['NOTE_CONTACT_DATE'] if dataset=='notes' else schema.get(date_field,[])
    date=_field_index(names,[overrides[date_field]] if date_field in overrides else date_alias)
    flag=_field_index(names,[overrides['anchor_flag']] if 'anchor_flag' in overrides else ['Anchor']) if dataset=='notes' else None
    note_type=_field_index(names,[overrides['note_type']] if 'note_type' in overrides else ['NOTE_TYPE']) if dataset=='notes' else None
    note_fields={}
    if dataset=='notes':
        for key,aliases in {'note_id':['NOTE_ID'],'note_csn':['NOTE_CSN'],
                            'encounter':['PAT_ENC_CSN_ID'],'line':['LINE']}.items():
            note_fields[key]=_field_index(names,[overrides[key]] if key in overrides else aliases)
    return {'headers':names,'patient':patient,'mrn':primary,'alternatives':alternatives,
            'date':date,'flag':flag,'note_type':note_type,'note_fields':note_fields,'status':status}


def _workbook(root,dataset,config):
    given=config.get('files',{}).get(dataset)
    stem='Notes' if dataset=='notes' else STEMS[dataset]
    candidates=[root/given] if given else [p for p in root.iterdir() if p.is_file() and
        not p.name.startswith('~$') and normalize(p.stem)==normalize(stem) and
        p.suffix.lower() in {'.xlsx','.xlsm','.xls'}]
    if len(candidates)!=1 or not candidates[0].is_file():
        raise PipelineError('Missing or ambiguous workbook for '+dataset+'; configure files.'+dataset+'.')
    path=candidates[0]
    if path.resolve().parent != root.resolve(): raise PipelineError('Configured workbook outside data-dir.')
    return path


@contextmanager
def _rows(path,dataset,config):
    sheet=config.get('sheets',{}).get(dataset)
    header_row=config.get('header_rows',{}).get(dataset,0)
    if not isinstance(header_row,int) or header_row<0: raise PipelineError('Invalid header_rows.'+dataset+'.')
    if path.suffix.lower()=='.xls':
        book=xlrd.open_workbook(str(path),on_demand=True)
        try:
            if sheet is None:
                if book.nsheets!=1: raise PipelineError('Multiple sheets in '+dataset+'; configure sheets.'+dataset+'.')
                ws=book.sheet_by_index(0)
            else: ws=book.sheet_by_name(sheet)
            if header_row>=ws.nrows: raise PipelineError('Header row missing in '+dataset+'.')
            headers=ws.row_values(header_row)
            def stream():
                for i in range(header_row+1,ws.nrows):
                    yield ws.row_values(i)
            yield ws.name,header_row,headers,stream()
        finally: book.release_resources()
    else:
        book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
        try:
            if sheet is None:
                if len(book.sheetnames)!=1: raise PipelineError('Multiple sheets in '+dataset+'; configure sheets.'+dataset+'.')
                sheet=book.sheetnames[0]
            if sheet not in book.sheetnames: raise PipelineError('Configured sheet missing in '+dataset+'.')
            ws=book[sheet]
            ws.reset_dimensions()
            stream=ws.iter_rows(values_only=True)
            for _ in range(header_row): next(stream,None)
            headers=next(stream,None)
            if headers is None: raise PipelineError('Header row missing in '+dataset+'.')
            yield sheet,header_row,headers,stream
        finally: book.close()


def _date_values(raw,config):
    parsed,failures=parse_dates(pd.Series(raw,dtype=object),config)
    return [(None,None) if pd.isna(v) else (v.isoformat(),v.strftime('%Y-%m-%d')) for v in parsed],failures


def _read_dataset(db,root,dataset,config,report):
    path=_workbook(root,dataset,config)
    with _rows(path,dataset,config) as (sheet,header_row,headers,stream):
        mapping=_discover(headers,dataset,config)
        status=mapping['status']
        stats=Counter()
        batch=[]
        batch_size=config.get('mrn_audit',{}).get('batch_rows',2000)
        if not isinstance(batch_size,int) or batch_size<1: raise PipelineError('mrn_audit.batch_rows must be a positive integer.')
        def flush():
            if not batch: return
            dates,failures=_date_values([item[4] for item in batch],config)
            stats['invalid_date_rows']+=failures
            def encoded_key(entry, parsed):
                parts=entry[6]
                if parts is None:return None
                if 'date' in configured_keys:
                    parts[configured_keys.index('date')]=parsed
                return json.dumps(parts,ensure_ascii=False) if all(part is not None for part in parts) else None
            db.executemany('''INSERT INTO audit_rows
              (dataset,ordinal,pat_id,mrn,mrn_status,date,day,note_flag,note_key,note_type)
              VALUES (?,?,?,?,?,?,?,?,?,?)''',[
                (dataset,entry[0],entry[1],entry[2],entry[3],dates[i][0],dates[i][1],entry[5],encoded_key(entry,dates[i][0]),entry[7])
                for i,entry in enumerate(batch)])
            db.commit();batch.clear()
        def value(values,index): return values[index] if index is not None and index<len(values) else None
        settings=config.get('notes',{}) if dataset=='notes' else {}
        note_fields=mapping['note_fields']
        configured_keys=settings.get('note_key_fields',[]) if dataset=='notes' else []
        candidate_key_available=bool(configured_keys and
                             all(note_fields.get(field) is not None or
                                 (field=='patient' and mapping['patient'] is not None) or
                                 (field=='date' and mapping['date'] is not None) for field in configured_keys))
        note_key_usable=bool(settings.get('note_key_verified') and settings.get('note_key_fields') and candidate_key_available)
        trailing=0
        for row in stream:
            values=list(row)
            if len(values)>len(headers) and any(missing_value(v) is not None for v in values[len(headers):]):
                raise PipelineError('Populated cells beyond configured header width in '+dataset+'.')
            values=(values+[None]*len(headers))[:len(headers)]
            if all(missing_value(v) is None for v in values):
                trailing+=1;continue
            for _ in range(trailing):
                stats['total_rows']+=1;stats['missing_mrn_rows']+=1
                batch.append((stats['total_rows']-1,None,None,'missing',None,'Missing',None,None))
                if len(batch)>=batch_size:flush()
            trailing=0
            stats['total_rows']+=1
            pat=_clean_id(value(values,mapping['patient']))
            primary=value(values,mapping['mrn'])
            mrn,mrn_status,trimmed=_clean_mrn(primary)
            if trimmed:stats['trimmed_mrn_rows']+=1
            if mrn is None:stats['missing_mrn_rows']+=1
            if mrn_status=='numeric_or_nontext_cell':stats['numeric_mrn_cell_rows']+=1
            for idx in mapping['alternatives']:
                alt,alt_status,_=_clean_mrn(value(values,idx))
                if alt_status=='numeric_or_nontext_cell':stats['numeric_alternative_mrn_rows']+=1
                if mrn is None and alt is not None: stats['alternate_present_preferred_missing_rows']+=1
                elif mrn is not None and alt is not None and alt!=mrn:
                    stats['conflicting_mrn_columns_rows']+=1
                    mrn=None;mrn_status='conflicting_mrn_columns';stats['missing_mrn_rows']+=1
                    break
            if mapping['mrn'] is None: mrn_status=status
            raw_date=value(values,mapping['date'])
            raw_date=missing_value(raw_date)
            flag=anchor_status(missing_value(value(values,mapping['flag']))) if dataset=='notes' else None
            note_type=missing_value(value(values,mapping['note_type'])) if dataset=='notes' else None
            note_key=None
            if dataset=='notes' and candidate_key_available:
                parts=[]
                for field in configured_keys:
                    idx=mapping['patient'] if field=='patient' else mapping['date'] if field=='date' else note_fields.get(field)
                    component=_clean_id(value(values,idx)) if field in {'patient','note_id','note_csn','encounter'} else missing_value(value(values,idx))
                    parts.append(component)
                note_key=parts
            batch.append((stats['total_rows']-1,pat,mrn,mrn_status,raw_date,flag,note_key,note_type))
            if len(batch)>=batch_size:flush()
        flush()
    stats['trailing_blank_rows_ignored']=trailing
    stats['note_key_usable']=int(note_key_usable)
    stats['candidate_note_key_available']=int(candidate_key_available)
    stats['mapped_pat_id']=int(mapping['patient'] is not None)
    stats['mapped_mrn']=int(mapping['mrn'] is not None)
    stats['mapped_date']=int(mapping['date'] is not None)
    stats['mapped_anchor_flag']=int(mapping['flag'] is not None)
    stats['file']=path.name;stats['sheet']=sheet;stats['header_row']=header_row;stats['mapping']=mapping
    report.data['completed_modules'].append('mrn_scan_'+dataset)
    return stats


def _stage(db):
    db.execute('PRAGMA cache_size=-16384')
    db.execute('PRAGMA temp_store=FILE')
    db.execute('''CREATE TABLE audit_rows (dataset TEXT,ordinal INTEGER,pat_id TEXT,mrn TEXT,
      mrn_status TEXT,date TEXT,day TEXT,note_flag TEXT,note_key TEXT,note_type TEXT)''')
    db.execute('CREATE INDEX patient_dataset ON audit_rows(dataset,mrn,pat_id)')
    db.execute('CREATE INDEX patient_pat_id ON audit_rows(dataset,pat_id)')
    db.execute('CREATE INDEX note_key_index ON audit_rows(note_key)')


def _sets(db,datasets,key):
    return {dataset:{row[0] for row in db.execute(f'SELECT DISTINCT {key} FROM audit_rows WHERE dataset=? AND {key} IS NOT NULL',(dataset,))}
            for dataset in datasets}


def _usable_mrn(info):
    return bool(info and info.get('mapped_mrn') and
                (info['total_rows']==0 or info.get('unique_mrns',0)>0))


def _usable_pat(info):
    return bool(info and info.get('mapped_pat_id') and
                (info['total_rows']==0 or info.get('unique_pat_ids',0)>0))


def _mapping_stats(db,dataset,stats):
    pat_set={row[0] for row in db.execute('SELECT DISTINCT pat_id FROM audit_rows WHERE dataset=? AND pat_id IS NOT NULL',(dataset,))}
    mrn_set={row[0] for row in db.execute('SELECT DISTINCT mrn FROM audit_rows WHERE dataset=? AND mrn IS NOT NULL',(dataset,))}
    missing_pat={row[0] for row in db.execute('SELECT DISTINCT pat_id FROM audit_rows WHERE dataset=? AND mrn IS NULL AND pat_id IS NOT NULL',(dataset,))}
    many_pat=db.execute('''SELECT COUNT(*) FROM
       (SELECT mrn FROM audit_rows WHERE dataset=? AND pat_id IS NOT NULL AND mrn IS NOT NULL
        GROUP BY mrn HAVING COUNT(DISTINCT pat_id)>1)''',(dataset,)).fetchone()[0]
    many_mrn=db.execute('''SELECT COUNT(*) FROM
       (SELECT pat_id FROM audit_rows WHERE dataset=? AND pat_id IS NOT NULL AND mrn IS NOT NULL
        GROUP BY pat_id HAVING COUNT(DISTINCT mrn)>1)''',(dataset,)).fetchone()[0]
    all_missing=db.execute('''SELECT COUNT(*) FROM
       (SELECT pat_id FROM audit_rows WHERE dataset=? AND pat_id IS NOT NULL
        GROUP BY pat_id HAVING SUM(CASE WHEN mrn IS NOT NULL THEN 1 ELSE 0 END)=0)''',(dataset,)).fetchone()[0]
    no_both=db.execute('SELECT COUNT(*) FROM audit_rows WHERE dataset=? AND pat_id IS NULL AND mrn IS NULL',(dataset,)).fetchone()[0]
    folded=Counter(value.casefold() for value in mrn_set)
    variants=sum(count>1 for count in folded.values())
    return {'unique_pat_ids':len(pat_set),'unique_mrns':len(mrn_set),
            'pat_ids_with_missing_mrn':len(missing_pat),
            'pat_ids_all_rows_missing_mrn':all_missing,'rows_missing_both_ids':no_both,
            'mrns_with_multiple_pat_ids':many_pat,'pat_ids_with_multiple_mrns':many_mrn,
            'mrn_values_with_case_variants':variants}


def _aggregate_groups(db,pat_sets,mrn_sets,stats,out):
    demos_pat=pat_sets.get('demographics') if _usable_pat(stats.get('demographics')) else None
    demos_mrn=mrn_sets.get('demographics') if _usable_mrn(stats.get('demographics')) else None
    inventory=[];mappings=[];intersections=[];comparison=[];schema=[];coverage=[];pairwise=[]
    for dataset in DATASETS:
        info=stats.get(dataset)
        if info is None:
            reason='dataset_unavailable'
            schema.append([dataset,*[None]*8,reason,*[None]*8,reason])
            inventory.append([dataset,*[None]*7,reason])
            mappings.append([dataset,*[None]*12,reason])
            intersections.append([dataset,None,None,None,None,None,None,None,reason])
            comparison.append([dataset,None,None,None,None,None,None,None,None,reason])
            continue
        m=info['mapping'];p=pat_sets.get(dataset) if _usable_pat(info) else None
        r=mrn_sets.get(dataset) if _usable_mrn(info) else None
        p_overlap=len(p&demos_pat) if p is not None and demos_pat is not None else None
        r_overlap=len(r&demos_mrn) if r is not None and demos_mrn is not None else None
        pfrac=p_overlap/len(demos_pat) if p_overlap is not None and demos_pat else None
        rfrac=r_overlap/len(demos_mrn) if r_overlap is not None and demos_mrn else None
        status='available' if r is not None and demos_mrn is not None else 'mrn_linkage_unavailable'
        if info['conflicting_mrn_columns_rows'] or info['numeric_mrn_cell_rows'] or info['missing_mrn_rows']:
            status = status+'_partial' if status=='available' else status
        if p is None and status.startswith('available'):
            status += '_pat_id_comparison_unavailable'
        schema.append([dataset,info['file'],info['sheet'],info['header_row'],json.dumps(m['headers'],ensure_ascii=False),
            m['headers'][m['patient']] if m['patient'] is not None else None,
            m['headers'][m['mrn']] if m['mrn'] is not None else None,
            json.dumps([m['headers'][i] for i in m['alternatives']],ensure_ascii=False),
            m['headers'][m['date']] if m['date'] is not None else None,m['status'],
            info['total_rows'],info['missing_mrn_rows'],info['numeric_mrn_cell_rows'],
            info['numeric_alternative_mrn_rows'],
            info['conflicting_mrn_columns_rows'],info['alternate_present_preferred_missing_rows'],
            info['trimmed_mrn_rows'],info['invalid_date_rows'],None])
        inventory.append([dataset,info['total_rows'],info.get('duplicate_rows'),len(p) if p is not None else None,
            len(r) if r is not None else None,info['missing_mrn_rows'],
            info['pat_ids_with_missing_mrn'] if p is not None else None,
            info['pat_ids_all_rows_missing_mrn'] if p is not None else None,status])
        mappings.append([dataset,info['total_rows'],len(p) if p is not None else None,
            len(r) if r is not None else None,info['missing_mrn_rows'],
            info['pat_ids_with_missing_mrn'] if p is not None else None,
            info['pat_ids_all_rows_missing_mrn'] if p is not None else None,
            info['rows_missing_both_ids'] if p is not None else None,
            info['mrns_with_multiple_pat_ids'] if p is not None else None,
            info['pat_ids_with_multiple_mrns'] if p is not None else None,
            info['mrn_values_with_case_variants'],info['numeric_mrn_cell_rows'],
            info['conflicting_mrn_columns_rows'],status])
        intersections.append([dataset,len(p) if p is not None else None,len(r) if r is not None else None,
            p_overlap,r_overlap,len(r-demos_mrn) if r is not None and demos_mrn is not None else None,
            len(demos_mrn-r) if r is not None and demos_mrn is not None else None,rfrac,status])
        comparison.append([dataset,info['total_rows'],len(p) if p is not None else None,
            len(r) if r is not None else None,len(r)-len(p) if p is not None and r is not None else None,
            p_overlap,r_overlap,pfrac,rfrac,status])
        if dataset!='demographics':
            coverage.append(['demographics_reference','demographics',dataset,len(demos_pat) if demos_pat is not None else None,
                len(demos_mrn) if demos_mrn is not None else None,p_overlap,r_overlap,pfrac,rfrac,status])
    for i,a in enumerate(DATASETS):
        for b in DATASETS[i+1:]:
            pa,pb=pat_sets.get(a),pat_sets.get(b)
            ra,rb=mrn_sets.get(a),mrn_sets.get(b)
            po=len(pa&pb) if pa is not None and pb is not None else None
            ro=len(ra&rb) if ra is not None and rb is not None else None
            pj=po/len(pa|pb) if po is not None and (pa|pb) else None
            rj=ro/len(ra|rb) if ro is not None and (ra|rb) else None
            case_only=(len({v.casefold() for v in ra}&{v.casefold() for v in rb})-
                       len({v.casefold() for v in ra&rb})) if ro is not None else None
            state='available' if ro is not None else 'mrn_linkage_unavailable'
            pairwise.append([a,b,po,ro,pj,rj,case_only,state])
            coverage.append(['pair',a,b,len(pa) if pa is not None else None,len(ra) if ra is not None else None,
                po,ro,pj,rj,state])
    complete_pat=all(dataset in pat_sets for dataset in DATASETS)
    complete_mrn=all(dataset in mrn_sets for dataset in DATASETS)
    all_pat=len(set.intersection(*(pat_sets[d] for d in DATASETS))) if complete_pat else None
    all_mrn=len(set.intersection(*(mrn_sets[d] for d in DATASETS))) if complete_mrn else None
    coverage.append(['all_eight',None,None,None,None,all_pat,all_mrn,None,None,
                     'available' if complete_mrn else 'incomplete_eight_dataset_linkage'])
    for name,rows in [('mrn_schema_validation.csv',schema),('mrn_patient_mapping_audit.csv',mappings),
                      ('mrn_dataset_inventory.csv',inventory),('mrn_patient_intersections.csv',intersections),
                      ('mrn_vs_pat_id_comparison.csv',comparison),('mrn_coverage_comparison.csv',coverage),
                      ('mrn_pairwise_intersections.csv',pairwise)]:write_csv(out/name,OUTPUTS[name],rows)
    complete_pairs=[row for row in pairwise if row[4] is not None and row[5] is not None]
    return {'demographics_pat_ids':len(demos_pat) if demos_pat is not None else None,
            'demographics_mrns':len(demos_mrn) if demos_mrn is not None else None,
            'all_eight_pat_ids':all_pat,'all_eight_mrns':all_mrn,
            'all_eight_available':complete_mrn,
            'exact_eight_dataset_mrn_equality':(all(mrn_sets[d]==demos_mrn for d in DATASETS) if complete_mrn else None),
            'pairwise_jaccard_improved':sum(row[5]>row[4] for row in complete_pairs),
            'pairwise_jaccard_decreased':sum(row[5]<row[4] for row in complete_pairs),
            'pairwise_jaccard_unchanged':sum(row[5]==row[4] for row in complete_pairs),
            'pairwise_comparisons_available':len(complete_pairs),
            'potential_case_only_mrn_matches_across_pairs':sum(row[6] for row in pairwise if row[6] is not None),
            'mapping_issues':sum(info['mrns_with_multiple_pat_ids']+info['pat_ids_with_multiple_mrns']+
                info['numeric_mrn_cell_rows']+info['conflicting_mrn_columns_rows'] for info in stats.values())}


def _demographic_anchors(db,out):
    """Compare original distinct calendar dates by PAT_ID and direct MRN."""
    result={};detail=[]
    for column in ['pat_id','mrn']:
        patient_days=defaultdict(set);patient_timestamps=defaultdict(set);patient_pats=defaultdict(set)
        patient_invalid=set()
        for pat,mrn,date,day in db.execute('SELECT pat_id,mrn,date,day FROM audit_rows WHERE dataset="demographics"'):
            person=pat if column=='pat_id' else mrn
            if person is None:continue
            if day is None:patient_invalid.add(person)
            else:patient_days[person].add(day);patient_timestamps[person].add(date)
            if pat is not None:patient_pats[person].add(pat)
        all_people={row[0] for row in db.execute('SELECT DISTINCT '+column+' FROM audit_rows WHERE dataset="demographics" AND '+column+' IS NOT NULL')}
        ambiguous={person for person in all_people if len(patient_days[person])>1}
        result[column]={'patients':len(all_people),'unique_day_patients':sum(len(patient_days[p])==1 and p not in patient_invalid for p in all_people),
            'multiple_day_patients':len(ambiguous),
            'no_valid_anchor_patients':sum(not patient_days[p] for p in all_people),
            'same_day_multiple_timestamp_patients':sum(len(patient_days[p])==1 and len(patient_timestamps[p])>1 for p in all_people),
            'patients_with_invalid_anchor_rows':len(patient_invalid),
            'ambiguous':ambiguous,'days':patient_days,'timestamps':patient_timestamps,
            'pat_ids':patient_pats,'all':all_people,'invalid':patient_invalid}
    pat_ambiguous=result['pat_id']['ambiguous'];mrn_ambiguous=result['mrn']['ambiguous']
    pat_to_mrn=defaultdict(set)
    pat_with_missing_mrn=set()
    for pat,mrn in db.execute('SELECT DISTINCT pat_id,mrn FROM audit_rows WHERE dataset="demographics" AND pat_id IS NOT NULL'):
        if mrn is not None:pat_to_mrn[pat].add(mrn)
        else:pat_with_missing_mrn.add(pat)
    previously_ambiguous_with_one_mrn={p for p in pat_ambiguous if len(pat_to_mrn[p])==1}
    previous_mrns={next(iter(pat_to_mrn[p])) for p in previously_ambiguous_with_one_mrn}
    multiple={
        'original_pat_ids_with_multiple_anchor_days':len(pat_ambiguous),
        'mrns_with_multiple_anchor_days':len(mrn_ambiguous),
        'original_ambiguous_pat_ids_with_one_mrn':len(previously_ambiguous_with_one_mrn),
        'original_ambiguous_pat_ids_missing_or_multiple_mrns':len(pat_ambiguous)-len(previously_ambiguous_with_one_mrn),
        'original_ambiguous_pat_ids_with_any_missing_mrn_rows':len(pat_ambiguous&pat_with_missing_mrn),
        'original_ambiguous_pat_ids_without_any_mrn':sum(not pat_to_mrn[p] for p in pat_ambiguous),
        'original_ambiguous_pat_ids_with_multiple_mrns':sum(len(pat_to_mrn[p])>1 for p in pat_ambiguous),
        'distinct_mrns_from_original_ambiguous_pat_ids':len(previous_mrns),
        'original_ambiguous_pat_ids_sharing_mrn_with_another_original_ambiguous':sum(
            any(len({q for q in pat_ambiguous if mrn in pat_to_mrn[q]})>1 for mrn in pat_to_mrn[p])
            for p in pat_ambiguous),
        'those_mrns_still_multiple_day':len(previous_mrns&mrn_ambiguous),
        'those_mrns_now_one_or_no_day':len(previous_mrns-mrn_ambiguous),
        'new_ambiguous_mrns_not_linked_to_old_ambiguous_pat_ids':len(mrn_ambiguous-previous_mrns),
        'new_ambiguous_mrns_from_multiple_individually_unambiguous_pat_ids':sum(
            len(result['mrn']['pat_ids'][p])>1 and not (result['mrn']['pat_ids'][p]&pat_ambiguous)
            for p in mrn_ambiguous),
        'ambiguous_mrns_with_multiple_pat_ids':sum(len(result['mrn']['pat_ids'][p])>1 for p in mrn_ambiguous),
        'all_demographic_mrns_with_multiple_pat_ids':sum(len(pats)>1 for pats in result['mrn']['pat_ids'].values()),
    }
    with (out/'mrn_anchor_patient_detail_SERDA_ONLY.csv').open('w',newline='',encoding='utf-8') as handle:
        writer=csv.writer(handle);writer.writerow(['MRN','distinct_anchor_days','PAT_ID_count',
            'contains_previous_multiple_anchor_PAT_ID','anchor_days'])
        for mrn in sorted(result['mrn']['all']):
            writer.writerow([mrn,len(result['mrn']['days'][mrn]),len(result['mrn']['pat_ids'][mrn]),
                bool(result['mrn']['pat_ids'][mrn]&pat_ambiguous),
                ';'.join(sorted(result['mrn']['days'][mrn]))])
    return result,multiple


def _mapping_detail(db,out):
    with (out/'mrn_patient_mapping_detail_SERDA_ONLY.csv').open('w',newline='',encoding='utf-8') as handle:
        writer=csv.writer(handle);writer.writerow(['dataset','PAT_ID','MRN','row_count'])
        for row in db.execute('''SELECT dataset,pat_id,mrn,COUNT(*) FROM audit_rows
                WHERE pat_id IS NOT NULL OR mrn IS NOT NULL GROUP BY dataset,pat_id,mrn'''):
            writer.writerow(row)


def _mapping_distributions(db,out,stats):
    specs=[('mrn_pat_ids_per_mrn_distribution.csv','mrn','pat_id'),
           ('mrn_mrns_per_pat_id_distribution.csv','pat_id','mrn')]
    for filename,person,partner in specs:
        rows=[]
        for dataset in DATASETS:
            if dataset not in stats:
                rows.append([dataset,None,None]);continue
            counts=[row[0] for row in db.execute(f'''SELECT COUNT(DISTINCT {partner}) FROM audit_rows
                WHERE dataset=? AND {person} IS NOT NULL GROUP BY {person}''',(dataset,))]
            rows.extend([dataset,count,n] for count,n in sorted(Counter(counts).items()))
        counts=[row[0] for row in db.execute(f'''SELECT COUNT(DISTINCT {partner}) FROM audit_rows
            WHERE {person} IS NOT NULL GROUP BY {person}''')]
        rows.extend(['ALL_DATASETS',count,n] for count,n in sorted(Counter(counts).items()))
        write_csv(out/filename,OUTPUTS[filename],rows)
    rows=[]
    for dataset in DATASETS:
        for key in ['pat_id','mrn']:
            if dataset not in stats:
                rows.append([dataset,key,None,None,None,None,'dataset_unavailable']);continue
            counts=[row[0] for row in db.execute(f'''SELECT COUNT(*) FROM audit_rows
                WHERE dataset=? AND {key} IS NOT NULL GROUP BY {key}''',(dataset,))]
            ordered=sorted(counts)
            rows.append([dataset,key,len(ordered),sum(c>1 for c in ordered),
                         ordered[(len(ordered)-1)//2] if ordered else None,
                         ordered[-1] if ordered else None,
                         'available' if ordered else 'no_linked_patients'])
    write_csv(out/'mrn_records_per_patient_summary.csv',OUTPUTS['mrn_records_per_patient_summary.csv'],rows)


def _notes_anchor(db,anchors,stats,out):
    """Row-date alignment under each identifier. No index visit is selected."""
    usable=stats.get('notes',{}).get('mapped_date') and stats.get('notes',{}).get('mapped_anchor_flag')
    if not usable:
        rows=[['anchor_note_fields_available',None,None,'unavailable_note_date_or_flag']]
        write_csv(out/'mrn_anchor_alignment_summary.csv',OUTPUTS['mrn_anchor_alignment_summary.csv'],rows)
        return {},{}
    result={};subgroup={}
    mrn_available=bool(_usable_mrn(stats.get('notes')) and _usable_mrn(stats.get('demographics')))
    pat_available=bool(_usable_pat(stats.get('notes')) and _usable_pat(stats.get('demographics')))
    for key in ['pat_id','mrn']:
        reference=anchors[key]
        yes_days=defaultdict(set);yes_patients=set();undated=set();candidate=defaultdict(set)
        for person,day,note_key,flag in db.execute('SELECT '+key+',day,note_key,note_flag FROM audit_rows WHERE dataset="notes"'):
            if person is None or flag!='Yes':continue
            yes_patients.add(person)
            if day:yes_days[person].add(day)
            else:undated.add(person)
            if note_key is not None:candidate[person].add(note_key)
        allcounts=Counter();multicounts=Counter()
        with (out/('mrn_notes_anchor_patient_detail_SERDA_ONLY.csv' if key=='mrn' else 'pat_id_notes_anchor_patient_detail_SERDA_ONLY.csv')).open('w',newline='',encoding='utf-8') as handle:
            writer=csv.writer(handle);writer.writerow([key.upper(),'reference_anchor_days','anchor_yes_note_days',
                'matching_days','candidate_anchor_note_keys','multiple_reference_anchor_days'])
            for person in sorted(reference['all']):
                days=yes_days[person];ref_days=reference['days'][person];matches=days&ref_days
                values={'patients_with_anchor_yes_rows':int(person in yes_patients),
                    'patients_with_matching_anchor_yes_day':int(bool(matches)),
                    'patients_without_matching_anchor_yes_day':int(not matches),
                    'patients_with_multiple_anchor_yes_dates':int(len(days)>1),
                    'patients_with_multiple_candidate_anchor_notes':int(len(candidate[person])>1),
                    'patients_with_undated_anchor_yes_rows':int(person in undated),
                    'patients_with_exact_anchor_day_set_agreement':int(bool(days) and days==ref_days and
                        person not in undated and person not in reference['invalid'])}
                allcounts.update(values)
                if person in reference['ambiguous']:multicounts.update(values)
                writer.writerow([person,len(ref_days),len(days),len(matches),len(candidate[person]),
                                 person in reference['ambiguous']])
        allcounts['patients_with_anchor_yes_rows']=len(yes_patients & reference['all'])
        allcounts['distinct_anchor_yes_row_dates']=sum(len(days) for days in yes_days.values())
        allcounts['note_identity_verified']=int(bool(stats['notes']['note_key_usable']))
        if not stats['notes']['candidate_note_key_available']:
            allcounts['patients_with_multiple_candidate_anchor_notes']=None
            multicounts['patients_with_multiple_candidate_anchor_notes']=None
        if (key=='mrn' and not mrn_available) or (key=='pat_id' and not pat_available):
            result[key]={metric:None for metric in allcounts}
            subgroup[key]={metric:None for metric in multicounts}
        else:
            result[key]=dict(allcounts);subgroup[key]=dict(multicounts)
    write_csv(out/'mrn_anchor_alignment_summary.csv',OUTPUTS['mrn_anchor_alignment_summary.csv'],[
        [metric,result['pat_id'].get(metric),result['mrn'].get(metric),'row_level_evidence'] for metric in sorted(set(result['pat_id'])|set(result['mrn']))])
    return result,subgroup


def _validated_note_counts(db,anchors,stats,out):
    """Only use note keys validated by the unchanged Notes reconstruction audit."""
    detail=out/'notes_reconstruction_patient_audit.csv'
    if not stats.get('notes',{}).get('note_key_usable') or not detail.exists():
        return None,'note_identity_unverified_or_reconstruction_unavailable'
    db.execute('CREATE TABLE verified_note_keys (note_key TEXT PRIMARY KEY)')
    with detail.open(newline='',encoding='utf-8') as handle:
        reader=csv.DictReader(handle)
        if not reader.fieldnames or 'candidate_key' not in reader.fieldnames or 'identity_validated' not in reader.fieldnames:
            return None,'notes_reconstruction_audit_unavailable'
        db.executemany('INSERT OR IGNORE INTO verified_note_keys VALUES (?)',
            ((row['candidate_key'],) for row in reader if row['identity_validated']=='True'))
    db.commit()
    yes={'pat_id':defaultdict(set),'mrn':defaultdict(set)}
    uncertain=Counter()
    for (key,pat_count,mrn_count,day_count,flag_count,flag,day,pat,mrn,
         missing_pat,missing_mrn,missing_day) in db.execute('''SELECT a.note_key,COUNT(DISTINCT a.pat_id),
            COUNT(DISTINCT a.mrn),COUNT(DISTINCT a.day),COUNT(DISTINCT a.note_flag),
            MIN(a.note_flag),MIN(a.day),MIN(a.pat_id),MIN(a.mrn),
            SUM(CASE WHEN a.pat_id IS NULL THEN 1 ELSE 0 END),
            SUM(CASE WHEN a.mrn IS NULL THEN 1 ELSE 0 END),
            SUM(CASE WHEN a.day IS NULL THEN 1 ELSE 0 END)
            FROM audit_rows a JOIN verified_note_keys v ON a.note_key=v.note_key
            WHERE a.dataset='notes' GROUP BY a.note_key'''):
        if flag_count!=1 or flag!='Yes':continue
        if day_count>1:uncertain['conflicting_note_dates']+=1;continue
        if missing_day:
            uncertain['undated_validated_note_groups']+=1
            day=None  # Note identity can be valid while date alignment remains unknown.
        if pat_count==1 and not missing_pat:yes['pat_id'][pat].add((key,day))
        else:uncertain['unlinked_pat_id_notes']+=1
        if mrn_count==1 and not missing_mrn:yes['mrn'][mrn].add((key,day))
        else:uncertain['unlinked_or_ambiguous_mrn_notes']+=1
    output={};subgroup={}
    for patient_key in ['pat_id','mrn']:
        reference=anchors[patient_key];source=yes[patient_key]
        linked={p for p in source if p in reference['all']}
        counts={
            'distinct_validated_anchor_yes_notes':sum(len(v) for v in source.values()),
            'demographics_patients_with_validated_anchor_yes_notes':len(linked),
            'demographics_patients_without_matching_validated_anchor_yes_note':sum(
                not any(day in reference['days'][p] for _,day in source.get(p,set())) for p in reference['all']),
            'demographics_patients_with_multiple_validated_anchor_yes_notes':sum(len(source.get(p,set()))>1 for p in reference['all']),
            'matching_validated_anchor_yes_notes':sum(day in reference['days'][p]
                for p in linked for _,day in source[p]),
        }
        multi={
            'multiple_anchor_patients_without_matching_validated_note':sum(
                not any(day in reference['days'][p] for _,day in source.get(p,set())) for p in reference['ambiguous']),
            'multiple_anchor_patients_with_multiple_validated_notes':sum(
                len(source.get(p,set()))>1 for p in reference['ambiguous']),
            'matching_validated_notes_for_multiple_anchor_patients':sum(day in reference['days'][p]
                for p in reference['ambiguous'] for _,day in source.get(p,set())),
        }
        output[patient_key]=counts;subgroup[patient_key]=multi
    output['uncertain_note_groups']=dict(uncertain)
    if not (_usable_pat(stats.get('notes')) and _usable_pat(stats.get('demographics'))):
        output['pat_id']={metric:None for metric in output['pat_id']}
        subgroup['pat_id']={metric:None for metric in subgroup['pat_id']}
    if not (_usable_mrn(stats.get('notes')) and _usable_mrn(stats.get('demographics'))):
        output['mrn']={metric:None for metric in output['mrn']}
        subgroup['mrn']={metric:None for metric in subgroup['mrn']}
    return (output,subgroup),'validated_note_keys_from_existing_notes_audit'


def _temporal(db,anchors,stats,config,out):
    reference=anchors['mrn'];rows=[]
    lookback=config.get('lookback_days',365);policy=config.get('anchor_day_policy','at_or_before_timestamp')
    for dataset in DATASETS:
        if dataset=='demographics':continue
        if dataset not in stats or not stats[dataset].get('mapped_mrn') or not stats[dataset].get('mapped_date'):
            rows.append([dataset,*[None]*7,'mrn_or_event_date_unavailable'])
            continue
        counts=Counter();status='available'
        for mrn,date,day in db.execute('SELECT mrn,date,day FROM audit_rows WHERE dataset=? AND mrn IS NOT NULL',(dataset,)):
            if mrn not in reference['all']:continue
            counts['matched_mrn_rows']+=1
            if len(reference['timestamps'][mrn])!=1 or mrn in reference['ambiguous'] or mrn in reference['invalid']:
                counts['ambiguous_reference_anchor_rows']+=1;continue
            anchor=pd.Timestamp(next(iter(reference['timestamps'][mrn])))
            if date is None:counts['missing_dates']+=1;continue
            event=pd.Timestamp(date)
            counts['dated_records']+=1
            if event>anchor:counts['post_anchor_records']+=1
            lower=anchor-pd.Timedelta(days=lookback)
            eligible=event>=lower and (event.normalize()<anchor.normalize() if policy=='pre_anchor_date_only' else event<=anchor)
            if eligible:counts['lookback_records']+=1
            elif day==anchor.strftime('%Y-%m-%d'):counts['anchor_day_excluded_rows']+=1
        rows.append([dataset,*[counts[k] for k in OUTPUTS['mrn_temporal_summary.csv'][1:-1]],status])
    write_csv(out/'mrn_temporal_summary.csv',OUTPUTS['mrn_temporal_summary.csv'],rows)


def _summary(report,out,aggregate,anchors,multiple,stats):
    original=aggregate.get('demographics_pat_ids');revised=aggregate.get('demographics_mrns')
    exact=aggregate.get('exact_eight_dataset_mrn_equality')
    empty_all=bool(stats) and all(info['total_rows']==0 for info in stats.values())
    any_missing=any(info['missing_mrn_rows']>0 for info in stats.values())
    statement=('No patient rows were observed; eight-dataset coverage cannot be assessed.' if empty_all else
        'All eight observed MRN sets exactly match; inspect missing MRN rows before interpreting completeness.' if exact is True else
        'Exact equality is not supported; approximate coverage has no agreed threshold.' if exact is False else
        'Eight-dataset MRN equality is unavailable because one or more mappings/files failed.')
    lines=['## MRN-Based Patient Linkage Validation','',
      'MRN audit uses direct validated MRN fields. PAT_ID-based counts are recomputed from the same workbooks for comparison.',
      f'Demographics unique PAT_ID: {original if original is not None else "unavailable"}; unique MRN: {revised if revised is not None else "unavailable"}.',
      f'Change after switching identifier: {revised-original if revised is not None and original is not None else "unavailable"}.',
      f'Previous reported checkpoint: 2,690 unique Demographics PAT_ID patients and 17 multiple-calendar-day PAT_ID cases; verify current PAT_ID recomputation against workbook version.',
      f'Patients present in all eight by PAT_ID: {aggregate.get("all_eight_pat_ids")}; by MRN: {aggregate.get("all_eight_mrns")}.',
      statement,
      ('Erika coverage expectation: observed exact equality supports it, pending local MRN-semantic validation.'
       if exact is True and not any_missing and not empty_all else
       'Erika coverage expectation: unresolved because MRNs are missing or an approximate-coverage threshold has not been agreed.'),
      f'Pairwise Jaccard overlap: {aggregate.get("pairwise_jaccard_improved")} improved, {aggregate.get("pairwise_jaccard_decreased")} decreased, {aggregate.get("pairwise_jaccard_unchanged")} unchanged of {aggregate.get("pairwise_comparisons_available")} comparable pairs.',
      f'Potential cross-dataset case-only MRN matches (diagnostic, never merged): {aggregate.get("potential_case_only_mrn_matches_across_pairs")}.',
      f'Unresolved mapping issue count across dataset summaries: {aggregate.get("mapping_issues")}.',
      f'PAT_ID values mapping to multiple MRNs across scanned datasets: {report.data.get("mrn_global_mapping",{}).get("pat_ids_with_multiple_mrns")}; MRNs mapping to multiple PAT_ID values: {report.data.get("mrn_global_mapping",{}).get("mrns_with_multiple_pat_ids")}.',
      f'Original multiple-calendar-day PAT_ID groups: {multiple.get("original_pat_ids_with_multiple_anchor_days")}; MRN groups: {multiple.get("mrns_with_multiple_anchor_days")}.',
      f'Old ambiguous PAT_ID groups linked to exactly one MRN: {multiple.get("original_ambiguous_pat_ids_with_one_mrn")}; associated MRNs still ambiguous: {multiple.get("those_mrns_still_multiple_day")}.',
      f'Old ambiguous PAT_ID groups sharing an MRN with another old ambiguous group: {multiple.get("original_ambiguous_pat_ids_sharing_mrn_with_another_original_ambiguous")}; with any missing MRN rows: {multiple.get("original_ambiguous_pat_ids_with_any_missing_mrn_rows")}.',
      f'New MRN ambiguous groups formed from multiple individually unambiguous PAT_ID groups: {multiple.get("new_ambiguous_mrns_from_multiple_individually_unambiguous_pat_ids")}.',
      'Anchor=Yes meaning and note identity require local clinical/export confirmation. No cohort membership, predictors, or labels changed.',
      'Missing or conflicting MRNs are excluded from MRN intersections, with no PAT_ID/name/encounter fallback.',
      'Detailed mappings and all identifier-bearing results remain inside SERDA. Clinical files have not been validated outside SERDA.']
    report.data['mrn_audit_summary']={k:v for k,v in aggregate.items() if k not in {'mrn_ids','pat_ids'}}
    report.data['mrn_audit_summary']['multiple_anchor_comparison']=multiple
    report.data['mrn_audit_lines']=lines


def audit_mrn(root,config,report,out,notes_only=False):
    """Audit eight exports (or Demographics+Notes); keep every output in a new run folder."""
    selected=['demographics','notes'] if notes_only else DATASETS
    stats={};failed={}
    try:
        with TemporaryDirectory(prefix='.mrn-stage-',dir=out) as temp:
            with closing(sqlite3.connect(str(Path(temp)/'mrn.sqlite'))) as db:
                db.execute("PRAGMA temp_store_directory='"+str(Path(temp).resolve()).replace("'","''")+"'")
                _stage(db)
                for dataset in selected:
                    try:
                        stats[dataset]=_read_dataset(db,root,dataset,config,report)
                    except Exception as exc:
                        message=str(exc) if isinstance(exc,PipelineError) else f'{type(exc).__name__}: MRN scan failed; review local workbook/configuration.'
                        failed[dataset]=message
                        report.fail('mrn_scan_'+dataset,message)
                # Half-scanned dataset rows cannot be presented as complete coverage.
                for dataset in failed:db.execute('DELETE FROM audit_rows WHERE dataset=?',(dataset,))
                db.commit()
                for dataset,info in stats.items():
                    original=report.data['datasets'].get(dataset,{})
                    info['duplicate_rows']=original.get('duplicate_rows')
                    info.update(_mapping_stats(db,dataset,info))
                    report.data.setdefault('mrn_datasets',{})[dataset]={
                        'row_count':info['total_rows'],'mrn_column':(
                            info['mapping']['headers'][info['mapping']['mrn']]
                            if info['mapping']['mrn'] is not None else None),
                        'mrn_mapping_status':info['mapping']['status'],
                        'unique_mrns':info['unique_mrns'] if _usable_mrn(info) else None,
                        'missing_mrn_rows':info['missing_mrn_rows'],
                        'numeric_mrn_cell_rows':info['numeric_mrn_cell_rows'],
                        'numeric_alternative_mrn_rows':info['numeric_alternative_mrn_rows'],
                        'conflicting_mrn_columns_rows':info['conflicting_mrn_columns_rows'],
                        'mrns_with_multiple_pat_ids':info['mrns_with_multiple_pat_ids'],
                        'pat_ids_with_multiple_mrns':info['pat_ids_with_multiple_mrns']}
                    if info['mapping']['mrn'] is None:
                        report.issue('mrn_schema_'+dataset,info['mapping']['status']+'; MRN patient counts unavailable.')
                _mapping_detail(db,out)
                _mapping_distributions(db,out,stats)
                pat_sets=_sets(db,[d for d in selected if _usable_pat(stats.get(d))],'pat_id')
                mrn_sets=_sets(db,[d for d in selected if _usable_mrn(stats.get(d))],'mrn')
                aggregate=_aggregate_groups(db,pat_sets,mrn_sets,stats,out)
                global_map={
                    'pat_ids_with_multiple_mrns':db.execute('''SELECT COUNT(*) FROM
                        (SELECT pat_id FROM audit_rows WHERE pat_id IS NOT NULL AND mrn IS NOT NULL
                         GROUP BY pat_id HAVING COUNT(DISTINCT mrn)>1)''').fetchone()[0],
                    'mrns_with_multiple_pat_ids':db.execute('''SELECT COUNT(*) FROM
                        (SELECT mrn FROM audit_rows WHERE pat_id IS NOT NULL AND mrn IS NOT NULL
                         GROUP BY mrn HAVING COUNT(DISTINCT pat_id)>1)''').fetchone()[0],
                }
                report.data['mrn_global_mapping']=global_map
                with (out/'mrn_patient_mapping_audit.csv').open('a',newline='',encoding='utf-8') as handle:
                    csv.writer(handle).writerow(['ALL_DATASETS',sum(info['total_rows'] for info in stats.values()),
                        db.execute('SELECT COUNT(DISTINCT pat_id) FROM audit_rows').fetchone()[0],
                        db.execute('SELECT COUNT(DISTINCT mrn) FROM audit_rows').fetchone()[0],
                        sum(info['missing_mrn_rows'] for info in stats.values()),
                        None,None,None,global_map['mrns_with_multiple_pat_ids'],global_map['pat_ids_with_multiple_mrns'],
                        None,sum(info['numeric_mrn_cell_rows'] for info in stats.values()),
                        sum(info['conflicting_mrn_columns_rows'] for info in stats.values()),'aggregate_across_scanned_datasets'])
                if 'demographics' in stats and (stats['demographics']['mapped_pat_id'] or stats['demographics']['mapped_mrn']):
                    anchors,multiple=_demographic_anchors(db,out)
                    align,subset=_notes_anchor(db,anchors,stats,out)
                    baseline=anchors['pat_id'];new=anchors['mrn']
                    fields=['patients','unique_day_patients','multiple_day_patients','no_valid_anchor_patients',
                            'same_day_multiple_timestamp_patients','patients_with_invalid_anchor_rows']
                    alignment_rows=[[field,baseline.get(field) if _usable_pat(stats['demographics']) else None,
                                       new.get(field) if _usable_mrn(stats['demographics']) else None,'demographics_anchor'] for field in fields]
                    if align:
                        alignment_rows += [[metric,align['pat_id'].get(metric),align['mrn'].get(metric),
                            'unverified_candidate_note_keys' if 'candidate_anchor_notes' in metric and
                            not stats.get('notes',{}).get('note_key_usable') else 'anchor_yes_rows']
                            for metric in sorted(set(align['pat_id'])|set(align['mrn']))]
                    write_csv(out/'mrn_anchor_alignment_summary.csv',OUTPUTS['mrn_anchor_alignment_summary.csv'],alignment_rows)
                    multi_rows=[[metric,None,(value if _usable_mrn(stats['demographics']) or
                        metric=='original_pat_ids_with_multiple_anchor_days' else None),
                        'mrn_transition' if _usable_mrn(stats['demographics']) else 'mrn_linkage_unavailable']
                        for metric,value in multiple.items()]
                    if subset:
                        multi_rows += [[metric,subset['pat_id'].get(metric),subset['mrn'].get(metric),'multiple_anchor_subgroup']
                            for metric in sorted(set(subset['pat_id'])|set(subset['mrn']))]
                    write_csv(out/'mrn_multiple_anchor_summary.csv',OUTPUTS['mrn_multiple_anchor_summary.csv'],multi_rows)
                    note_counts,note_status=_validated_note_counts(db,anchors,stats,out)
                    if note_counts:
                        note_values,note_subset=note_counts
                        with (out/'mrn_anchor_alignment_summary.csv').open('a',newline='',encoding='utf-8') as handle:
                            csv.writer(handle).writerows([[metric,note_values['pat_id'][metric],note_values['mrn'][metric],note_status]
                                for metric in note_values['pat_id']])
                        with (out/'mrn_multiple_anchor_summary.csv').open('a',newline='',encoding='utf-8') as handle:
                            csv.writer(handle).writerows([[metric,note_subset['pat_id'][metric],note_subset['mrn'][metric],note_status]
                                for metric in note_subset['pat_id']])
                        report.data['mrn_note_group_issues']=note_values['uncertain_note_groups']
                    else:
                        with (out/'mrn_anchor_alignment_summary.csv').open('a',newline='',encoding='utf-8') as handle:
                            csv.writer(handle).writerow(['distinct_validated_anchor_yes_notes',None,None,note_status])
                        report.skip('mrn_validated_note_counts',note_status)
                    if _usable_mrn(stats['demographics']):_temporal(db,anchors,stats,config,out)
                    report.data['completed_modules'].append('mrn_anchor_comparison')
                else:
                    anchors={};multiple={}
                    report.skip('mrn_anchor_comparison','Demographics identifier scan unavailable.')
                report.data['completed_modules'].extend(['mrn_schema_validation','mrn_patient_mapping','mrn_intersections'])
                _summary(report,out,aggregate,anchors,multiple,stats)
                if not aggregate['all_eight_available'] and not notes_only:
                    report.issue('mrn_coverage','All-eight MRN comparison unavailable until every dataset has a validated MRN column and successful scan.')
                if notes_only:report.skip('mrn_all_eight_coverage','--notes-audit-only scanned Demographics and Notes only.')
    except Exception as exc:
        reason=str(exc) if isinstance(exc,PipelineError) else f'{type(exc).__name__}: MRN audit failed; inspect local configuration and disk space.'
        report.fail('mrn_audit',reason)
    finally:
        report.data.setdefault('mrn_audit_lines',[
            '## MRN-Based Patient Linkage Validation','',
            'MRN audit unavailable or incomplete; review execution_report.json failed_modules and skipped_modules.'])
        for name,columns in OUTPUTS.items():
            if not (out/name).exists():write_csv(out/name,columns,[])
