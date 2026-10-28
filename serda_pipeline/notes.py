"""Streaming, disk-backed Notes audit. Run on clinical files only inside SERDA.

No Notes fields enter predictors or cohort selection. Text is never logged.
"""
from collections import Counter, defaultdict
from contextlib import ExitStack, closing
import csv
import hashlib
import json
from pathlib import Path
import re
import sqlite3
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import load_workbook

from .analysis import plot
from .core import PipelineError, parse_dates
from .schema import normalize

FIELDS = {
    'patient': ['PAT_ID'], 'encounter': ['PAT_ENC_CSN_ID'],
    'note_id': ['NOTE_ID'], 'note_csn': ['NOTE_CSN', 'NOTE_CSN_ID'], 'line': ['LINE'],
    'date': ['NOTE_CONTACT_DATE'], 'text': ['NOTE_TEXT'],
    'note_type': ['NOTE_TYPE'], 'anchor_flag': ['Anchor'],
}
MODULES = ['notes_inventory', 'notes_reconstruction', 'notes_patient_intersections',
           'notes_anchor_alignment']
FIGURES = ['notes_type_frequency', 'notes_anchor_rows', 'notes_reconstruction_quality',
           'notes_anchor_alignment', 'notes_per_patient']
DELIVERABLES = {
    'notes_inventory.csv': ['metric', 'value', 'status'],
    'notes_missingness.csv': ['column', 'missing_fraction', 'missing_rows', 'total_rows'],
    'notes_type_distribution.csv': ['note_type', 'row_count', 'validated_note_count'],
    'notes_reconstruction_quality.csv': ['metric', 'count', 'status'],
    'notes_anchor_summary.csv': ['level', 'anchor_status', 'count', 'status'],
    'notes_patient_intersections.csv': ['metric', 'count', 'status'],
    'notes_anchor_alignment_summary.csv': ['metric', 'count', 'status'],
    'notes_multiple_anchor_summary.csv': ['metric', 'count', 'status'],
}


def missing_value(value):
    """Same NULL/blank conventions as structured loading; preserve literal NA."""
    if value is None: return None
    text = str(value)
    return None if not text.strip() or text in {'NULL', 'null', '(null)', 'None'} else text


def identifier(value):
    if value is None: return None
    value = value.strip()
    if re.fullmatch(r'[0-9]+\.0+|[0-9.]+[eE][+-]?[0-9]+', value): return None
    return value


def anchor_status(value):
    if value is None: return 'Missing'
    value = value.strip().casefold()
    return {'yes': 'Yes', 'no': 'No'}.get(value, 'Other')


def line_number(value):
    if value is None: return None
    # LINE is an ordinal, so integral Excel numeric strings such as 1.0 are valid.
    if re.fullmatch(r'[0-9]+(?:\.0+)?', value.strip()):
        number = int(value.strip().split('.')[0])
        return number if number <= 9223372036854775807 else None
    return None


def line_quality(text_cardinalities, source_order, rows_with_line, start):
    """Ordinal/cardinality audit independent of any patient or clinical text."""
    positions = set(text_cardinalities)
    gaps = max(positions)-start+1-sum(n>=start for n in positions) if positions and max(positions)>=start else 0
    return {
        'duplicate_line_rows': rows_with_line-len(positions),
        'conflicting_lines': sum(count>1 for count in text_cardinalities.values()),
        'missing_line_positions': gaps,
        'unexpected_first_line': bool(positions) and min(positions)!=start,
        'no_valid_lines': not positions,
        'out_of_order': any(a>b for a,b in zip(source_order,source_order[1:])),
    }


def schema_mapping(headers, config):
    if not headers or any(value is None or not str(value).strip() for value in headers):
        raise PipelineError('Notes header has blank columns; verify header_rows.notes.')
    headers = [str(value) for value in headers]
    normalized = [normalize(value) for value in headers]
    if len(set(normalized)) != len(headers):
        raise PipelineError('Notes has duplicate/normalization-colliding headers.')
    mapping, ambiguous, absent = {}, [], []
    overrides = config.get('columns', {}).get('notes', {})
    for field, aliases in FIELDS.items():
        choices = [overrides[field]] if field in overrides else aliases
        matches = [i for i, value in enumerate(normalized) if value in {normalize(a) for a in choices}]
        if len(matches) == 1: mapping[field] = matches[0]
        elif len(matches) > 1: ambiguous.append(field)
        else: absent.append(field)
    return headers, mapping, ambiguous, absent


def validate_settings(config):
    settings = config.get('notes', {})
    keys = settings.get('note_key_fields', [])
    if not isinstance(keys, list) or len(keys) != len(set(keys)) or any(k not in FIELDS or k in {'line', 'text', 'anchor_flag'} for k in keys):
        raise PipelineError('notes.note_key_fields must be a unique list of mapped identifier/metadata fields, excluding line/text/anchor_flag.')
    if settings.get('note_key_verified', False) and not keys:
        raise PipelineError('Verified Notes identity requires explicit notes.note_key_fields.')
    for key, default in [('line_start', 1), ('max_fragments_per_note', 100000), ('max_note_characters', 5000000), ('batch_rows', 1000)]:
        value = settings.get(key, default)
        if not isinstance(value, int) or isinstance(value, bool) or value < (0 if key == 'line_start' else 1):
            raise PipelineError('Invalid notes.' + key)
    for key in ['note_key_verified', 'encounter_linkage_verified', 'export_reconstructed_text']:
        if key in settings and not isinstance(settings[key], bool):
            raise PipelineError('notes.' + key + ' must be true or false.')
    return settings


def write_csv(path, columns, rows):
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        writer.writerows(rows)


def metrics_file(out, name, values, status='available'):
    write_csv(out/name, DELIVERABLES[name], ((metric, value, status) for metric, value in values.items()))


def find_workbook(root, config, report):
    specified = config.get('files', {}).get('notes')
    candidates = [root/specified] if specified else [p for p in root.iterdir()
        if p.is_file() and not p.name.startswith('~$') and normalize(p.stem) == 'notes'
        and p.suffix.lower() in {'.xlsx', '.xlsm'}]
    if not candidates or not candidates[0].exists():
        report.data['missing_files'].append('notes')
        raise PipelineError('Notes workbook missing; configure files.notes with exact filename.')
    if len(candidates) != 1: raise PipelineError('Ambiguous Notes workbook matches; configure files.notes.')
    path = candidates[0]
    if path.resolve().parent != root.resolve() or path.suffix.lower() not in {'.xlsx', '.xlsm'}:
        raise PipelineError('Notes file must be an .xlsx/.xlsm workbook directly inside data-dir.')
    report.data['files_found']['notes'] = path.name
    return path


def prepare_reference(db, demo):
    """Use original Demographics, never the selected analysis cohort."""
    db.execute('CREATE TABLE reference (patient TEXT, day TEXT, PRIMARY KEY(patient,day))')
    db.execute('CREATE TABLE reference_patients (patient TEXT PRIMARY KEY)')
    db.execute('CREATE TABLE encounter_map (encounter TEXT PRIMARY KEY, patient TEXT)')
    if demo is None: return None
    patients = set(demo.patient.dropna())
    db.executemany('INSERT INTO reference_patients VALUES (?)', ((patient,) for patient in patients))
    dates = defaultdict(set)
    invalid = set()
    for patient, date in demo[['patient', 'anchor']].itertuples(index=False, name=None):
        if pd.isna(patient): continue
        if pd.isna(date): invalid.add(patient)
        else: dates[patient].add(date.strftime('%Y-%m-%d'))
    db.executemany('INSERT OR IGNORE INTO reference VALUES (?,?)',
                   ((p, day) for p, days in dates.items() for day in days))
    encounters = defaultdict(set)
    if 'encounter' in demo:
        for patient, encounter in demo[['patient', 'encounter']].itertuples(index=False, name=None):
            if pd.notna(patient) and pd.notna(encounter) and identifier(str(encounter)):
                encounters[identifier(str(encounter))].add(patient)
    db.executemany('INSERT INTO encounter_map VALUES (?,?)',
                   ((encounter, next(iter(ids)) if len(ids) == 1 else None) for encounter, ids in encounters.items()))
    return {'patients': patients, 'dates': dates, 'invalid_anchor_patients': invalid,
            'ambiguous': {p for p, days in dates.items() if len(days) > 1},
            'conflicting_encounters': sum(len(ids) > 1 for ids in encounters.values())}


def ingest(db, root, config, report, out):
    settings = validate_settings(config)
    path = find_workbook(root, config, report)
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        report.data.setdefault('workbook_sheets', {})['notes'] = book.sheetnames
        sheet = config.get('sheets', {}).get('notes')
        if sheet is None:
            if len(book.sheetnames) != 1: raise PipelineError('Multiple Notes sheets; configure sheets.notes.')
            sheet = book.sheetnames[0]
        if sheet not in book.sheetnames: raise PipelineError('Configured Notes sheet not found.')
        header_row = config.get('header_rows', {}).get('notes', 0)
        if not isinstance(header_row, int) or header_row < 0: raise PipelineError('Invalid header_rows.notes.')
        worksheet = book[sheet]
        # Ignore inaccurate exporter dimension metadata; stream actual XML rows.
        worksheet.reset_dimensions()
        stream = worksheet.iter_rows(values_only=True)
        for _ in range(header_row): next(stream, None)
        header = next(stream, ())
        headers, mapping, ambiguous, absent = schema_mapping(header, config)
        stats = {'real_columns': headers, 'column_mapping': {k: headers[v] for k, v in mapping.items()},
                 'missing_columns': absent, 'ambiguous_columns': ambiguous,
                 'unexpected_columns': [h for i, h in enumerate(headers) if i not in mapping.values()],
                 'row_count': 0, 'duplicate_rows': 0, 'missingness': {}, 'date_parsing_failures': {}}
        report.data['datasets']['notes'] = stats
        # Ambiguous fields are deliberately unmapped; raw inventory remains possible.
        if ambiguous: report.issue('notes_schema', 'Ambiguous mapped fields unavailable: ' + ', '.join(ambiguous))
        if absent: report.issue('notes_schema', 'Unmapped Notes fields (not assumed present): ' + ', '.join(absent))
        candidate_keys = settings.get('note_key_fields', [])  # Explicit export granularity; never prefer a column just because it exists.
        verified = bool(settings.get('note_key_verified', False))
        if any(key not in mapping for key in candidate_keys): verified = False
        if settings.get('note_key_verified') and not verified:
            report.fail('notes_identity', 'A configured verified key field is missing or ambiguous; reconstruction withheld.')
        if not candidate_keys: report.issue('notes_identity', 'No explicit notes.note_key_fields; raw row audit continues but no candidate text groups are assumed.')
        stats['note_key_fields'] = candidate_keys
        stats['note_identity_verified_by_local_configuration'] = verified
        stats['mapped_fields_missing_rows'] = Counter()
        db.execute('CREATE TABLE seen_rows (digest TEXT PRIMARY KEY, raw_json TEXT NOT NULL)')
        db.execute('''CREATE TABLE fragments (ordinal INTEGER PRIMARY KEY, key TEXT,
            raw_patient TEXT, encounter TEXT, patient TEXT, linkage_status TEXT,
            note_id TEXT, note_csn TEXT, line INTEGER, raw_line TEXT,
            date TEXT, day TEXT, raw_date TEXT, text TEXT, note_type TEXT, flag TEXT,
            exact_duplicate INTEGER)''')
        missing = [0] * len(headers)
        flag_counts, types = Counter(), Counter()
        raw_date_missing = invalid_dates = numeric_patient_ids = 0
        batch = []

        def flush():
            nonlocal invalid_dates
            if not batch: return
            parsed, failures = parse_dates(pd.Series([row['date'] for row in batch], dtype=object), config)
            invalid_dates += failures
            for entry, date in zip(batch, parsed):
                entry['parsed_date'] = None if pd.isna(date) else date.isoformat()
                entry['day'] = None if pd.isna(date) else date.strftime('%Y-%m-%d')
                # Date component of a configured key uses parsed timestamp, never invalid raw dates.
                key_values = [entry['parsed_date'] if k == 'date' else entry.get(k) for k in candidate_keys]
                key = json.dumps(key_values, ensure_ascii=False) if candidate_keys and all(v is not None for v in key_values) else None
                db.execute('INSERT INTO fragments VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (entry['ordinal'], key, entry['patient'], entry['encounter'], None, None,
                     entry['note_id'], entry['note_csn'], line_number(entry['line']), entry['line'],
                     entry['parsed_date'], entry['day'], entry['date'], entry['text'], entry['note_type'],
                     anchor_status(entry['anchor_flag']), entry['duplicate']))
            db.commit()
            batch.clear()

        def consume(values):
            nonlocal raw_date_missing, numeric_patient_ids
            stats['row_count'] += 1
            for i, value in enumerate(values): missing[i] += value is None
            raw_json = json.dumps(values, ensure_ascii=False, separators=(',', ':'))
            digest = hashlib.sha256(raw_json.encode('utf-8')).hexdigest()
            found = db.execute('SELECT raw_json FROM seen_rows WHERE digest=?', (digest,)).fetchone()
            if found and found[0] != raw_json: raise PipelineError('Exact duplicate hash collision; audit stopped.')
            duplicate = int(found is not None)
            if not duplicate: db.execute('INSERT INTO seen_rows VALUES (?,?)', (digest, raw_json))
            stats['duplicate_rows'] += duplicate
            entry = {field: values[mapping[field]] if field in mapping else None for field in FIELDS}
            for field, value in entry.items(): stats['mapped_fields_missing_rows'][field] += value is None
            if entry['patient'] is not None and identifier(entry['patient']) is None: numeric_patient_ids += 1
            entry['patient'] = identifier(entry['patient'])
            entry['encounter'] = identifier(entry['encounter'])
            entry['note_id'] = identifier(entry['note_id'])
            entry['note_csn'] = identifier(entry['note_csn'])
            entry['ordinal'] = stats['row_count'] - 1
            entry['duplicate'] = duplicate
            raw_date_missing += entry['date'] is None
            flag_counts[anchor_status(entry['anchor_flag'])] += 1
            types[entry['note_type'] or '(Missing)'] += 1
            batch.append(entry)
            if len(batch) >= settings.get('batch_rows', 1000): flush()

        pending_blank_rows = 0
        for row in stream:
            values = [missing_value(v) for v in row]
            if len(values) > len(headers) and any(v is not None for v in values[len(headers):]):
                raise PipelineError('Notes contains populated cells beyond the header width; correct header row.')
            values = (values + [None]*len(headers))[:len(headers)]
            if all(value is None for value in values):
                pending_blank_rows += 1
                continue
            for _ in range(pending_blank_rows): consume([None]*len(headers))
            pending_blank_rows = 0
            consume(values)
        # pandas.read_excel also discards trailing wholly empty rows.
        stats['trailing_blank_rows_ignored'] = pending_blank_rows
        flush()
        db.execute('CREATE INDEX fragment_key ON fragments(key,ordinal)')
        db.execute('CREATE INDEX fragment_csn ON fragments(note_csn)')
        stats['missingness'] = {h: missing[i]/stats['row_count'] if stats['row_count'] else None for i, h in enumerate(headers)}
        stats['date_parsing_failures']['date'] = invalid_dates
        stats['missing_note_date_rows'] = raw_date_missing
        stats['numeric_coerced_patient_id_rows'] = numeric_patient_ids
        write_csv(out/'notes_missingness.csv', DELIVERABLES['notes_missingness.csv'],
            ((h, stats['missingness'][h], missing[i], stats['row_count']) for i, h in enumerate(headers)))
        report.data['completed_modules'].append('load_notes')
        return stats, verified, types, flag_counts
    finally:
        book.close()


def establish_linkage(db, stats, settings, reference, report):
    mapping = stats['column_mapping']
    # NOTE_CSN is never silently equated with an encounter CSN.
    encounter_reliable = 'encounter' in mapping and (
        normalize(mapping['encounter']) == normalize('PAT_ENC_CSN_ID') or settings.get('encounter_linkage_verified', False))
    direct = 'patient' in mapping
    db.execute('''UPDATE fragments SET patient=raw_patient,
        linkage_status=CASE WHEN raw_patient IS NOT NULL THEN 'direct_patient' ELSE 'unlinked' END''')
    if encounter_reliable and reference is not None:
        db.execute('''UPDATE fragments SET patient=NULL, linkage_status='patient_encounter_conflict'
            WHERE raw_patient IS NOT NULL AND EXISTS (SELECT 1 FROM encounter_map e
                WHERE e.encounter=fragments.encounter AND e.patient IS NOT NULL AND e.patient!=fragments.raw_patient)''')
        db.execute('''UPDATE fragments SET patient=(SELECT patient FROM encounter_map e WHERE e.encounter=fragments.encounter),
            linkage_status='demographic_encounter' WHERE raw_patient IS NULL AND EXISTS
            (SELECT 1 FROM encounter_map e WHERE e.encounter=fragments.encounter AND e.patient IS NOT NULL)''')
    linked_patients = db.execute('SELECT COUNT(DISTINCT patient) FROM fragments').fetchone()[0]
    reliable = (direct or (encounter_reliable and reference is not None)) and (stats['row_count'] == 0 or linked_patients > 0)
    stats['patient_linkage_available'] = reliable
    stats['encounter_linkage_available'] = encounter_reliable and reference is not None
    stats['linkage_counts'] = dict(db.execute('SELECT linkage_status,COUNT(*) FROM fragments GROUP BY linkage_status'))
    stats['missing_patient_ids'] = db.execute('SELECT COUNT(*) FROM fragments WHERE patient IS NULL').fetchone()[0]
    stats['unique_patients'] = linked_patients if reliable else None
    stats['linked_rows'] = stats['row_count'] - stats['missing_patient_ids']
    if not reliable: report.issue('notes_linkage', 'No reliable mapped patient ID or validated encounter linkage. Patient counts/intersections/alignment unavailable; no MRN/name/NOTE_CSN fallback.')
    if reference: stats['reference_conflicting_encounters'] = reference['conflicting_encounters']
    db.commit()


def csn_granularity(db, stats, out):
    """Structural evidence, not proof that a CSN denotes one clinical note."""
    counts = Counter({key:0 for key in ['distinct_note_csns','csns_multiple_patients',
        'csns_multiple_note_ids','csns_multiple_calendar_dates','csns_multiple_note_types',
        'csns_repeated_line_numbers']})
    rows = db.execute('''SELECT note_csn,COUNT(*),COUNT(DISTINCT patient),COUNT(DISTINCT note_id),
        COUNT(DISTINCT day),COUNT(DISTINCT note_type),
        SUM(CASE WHEN line IS NOT NULL THEN 1 ELSE 0 END),COUNT(DISTINCT line) FROM fragments
        WHERE note_csn IS NOT NULL GROUP BY note_csn''')
    with (out/'notes_csn_granularity_patient_audit.csv').open('w',newline='',encoding='utf-8') as handle:
        writer=csv.writer(handle)
        writer.writerow(['NOTE_CSN','rows','patients','note_ids','calendar_dates','note_types','rows_with_valid_line','distinct_lines'])
        for row in rows:
            writer.writerow(row)
            counts['distinct_note_csns'] += 1
            counts['csns_repeated_line_numbers'] += row[6] > row[7]
            for i, field in [(2,'patients'),(3,'note_ids'),(4,'calendar_dates'),(5,'note_types')]:
                counts['csns_multiple_'+field] += row[i] > 1
    stats['note_csn_granularity'] = dict(counts)


def reconstruct(db, stats, verified, settings, report, out):
    db.execute('''CREATE TABLE notes (key TEXT PRIMARY KEY,patient TEXT,day TEXT,flag TEXT,
        note_type TEXT,structurally_complete INTEGER,validated INTEGER,reason TEXT)''')
    quality = Counter({key:0 for key in [
        'candidate_note_groups','structurally_complete_groups','failed_groups','validated_note_groups',
        'reconstructed_notes','groups_duplicate_lines','duplicate_line_rows','groups_conflicting_lines',
        'groups_missing_line_positions','missing_line_positions','groups_out_of_order_fragments',
        'groups_conflicting_anchor_flags','groups_unlinked_or_partially_linked','oversized_groups',
        'groups_missing_anchor_indicator','groups_missing_note_type']})
    quality['rows_without_candidate_key'] = db.execute('SELECT COUNT(*) FROM fragments WHERE key IS NULL').fetchone()[0]
    quality['rows_missing_line'] = db.execute('SELECT COUNT(*) FROM fragments WHERE raw_line IS NULL').fetchone()[0]
    quality['rows_invalid_line'] = db.execute('SELECT COUNT(*) FROM fragments WHERE raw_line IS NOT NULL AND line IS NULL').fetchone()[0]
    quality['rows_missing_text'] = db.execute('SELECT COUNT(*) FROM fragments WHERE text IS NULL').fetchone()[0]
    quality['rows_missing_date'] = stats['missing_note_date_rows']
    quality['rows_invalid_date'] = stats['date_parsing_failures']['date']
    quality['rows_missing_anchor_indicator'] = stats['mapped_fields_missing_rows']['anchor_flag']
    keys = db.execute('SELECT key,COUNT(*),COALESCE(SUM(LENGTH(text)),0) FROM fragments WHERE key IS NOT NULL GROUP BY key')
    with ExitStack() as stack:
        audit = stack.enter_context((out/'notes_reconstruction_patient_audit.csv').open('w',newline='',encoding='utf-8'))
        writer = csv.writer(audit)
        writer.writerow(['candidate_key','PAT_ID','rows','unique_fragments','duplicate_line_rows',
            'conflicting_lines','missing_line_positions','out_of_order','structurally_complete',
            'identity_validated','reconstruction_failure_reasons'])
        export = stack.enter_context((out/'notes_reconstructed_text_SERDA_only.jsonl').open('w',encoding='utf-8')) if settings.get('export_reconstructed_text') and verified else None
        for key, size, characters in keys:
            quality['candidate_note_groups'] += 1
            if size > settings.get('max_fragments_per_note',100000) or characters > settings.get('max_note_characters',5000000):
                reason = 'fragment_limit' if size > settings.get('max_fragments_per_note',100000) else 'text_size_limit'
                quality['oversized_groups'] += 1
                quality['failed_groups'] += 1
                quality['failure_'+reason] += 1
                db.execute('INSERT INTO notes VALUES (?,?,?,?,?,?,?,?)',(key,None,None,'Unresolved',None,0,0,reason))
                writer.writerow([key,None,size,None,None,None,None,None,False,False,reason])
                continue
            rows = db.execute('''SELECT patient,line,date,day,text,note_type,flag,exact_duplicate
                FROM fragments WHERE key=? ORDER BY ordinal''',(key,)).fetchall()
            patients = {r[0] for r in rows if r[0] is not None}
            dates = {r[2] for r in rows if r[2] is not None}
            days = {r[3] for r in rows if r[3] is not None}
            types = {r[5] for r in rows if r[5] is not None}
            flags = {r[6] for r in rows}
            reason = []
            if len(patients)>1: reason.append('multiple_patients')
            if len(dates)>1: reason.append('multiple_note_timestamps')
            if len(types)>1: reason.append('multiple_note_types')
            if len(flags)>1: reason.append('conflicting_anchor_flags')
            quality['groups_missing_anchor_indicator'] += 'Missing' in flags
            quality['groups_missing_note_type'] += any(r[5] is None for r in rows)
            if any(r[0] is None for r in rows): quality['groups_unlinked_or_partially_linked'] += 1
            if any(r[2] is None for r in rows): reason.append('missing_or_invalid_date')
            if any(r[4] is None for r in rows): reason.append('missing_text')
            if any(r[1] is None for r in rows): reason.append('missing_or_invalid_line')
            lines = defaultdict(set)
            ordered = []
            for row in rows:
                if row[1] is not None:
                    lines[row[1]].add(row[4])
                    if not row[7]: ordered.append(row[1])
            ordinal_quality = line_quality({n:len(texts) for n,texts in lines.items()},ordered,
                sum(r[1] is not None for r in rows),settings.get('line_start',1))
            duplicate_lines = ordinal_quality['duplicate_line_rows']
            conflicts = ordinal_quality['conflicting_lines']
            gaps = ordinal_quality['missing_line_positions']
            if conflicts: reason.append('conflicting_line_text')
            if ordinal_quality['unexpected_first_line']: reason.append('unexpected_first_line')
            if gaps: reason.append('missing_line_positions')
            if ordinal_quality['no_valid_lines']: reason.append('no_valid_lines')
            out_of_order = ordinal_quality['out_of_order']
            complete = not reason
            identity_validated = verified and len(patients)<=1 and len(dates)<=1 and len(types)<=1
            quality['structurally_complete_groups'] += complete
            quality['failed_groups'] += not complete
            quality['validated_note_groups'] += identity_validated
            quality['reconstructed_notes'] += complete and identity_validated
            quality['groups_duplicate_lines'] += duplicate_lines>0
            quality['duplicate_line_rows'] += duplicate_lines
            quality['groups_conflicting_lines'] += conflicts>0
            quality['groups_missing_line_positions'] += gaps>0
            quality['missing_line_positions'] += gaps
            quality['groups_out_of_order_fragments'] += out_of_order
            quality['groups_conflicting_anchor_flags'] += len(flags)>1
            for value in reason: quality['failure_'+value] += 1
            patient = next(iter(patients)) if len(patients)==1 and all(r[0] is not None for r in rows) else None
            day = next(iter(days)) if len(days)==1 and all(r[3] is not None for r in rows) else None
            flag = next(iter(flags)) if len(flags)==1 else 'Conflicting'
            note_type = next(iter(types)) if len(types)==1 else None
            db.execute('INSERT INTO notes VALUES (?,?,?,?,?,?,?,?)',
                (key,patient,day,flag,note_type,int(complete),int(identity_validated),';'.join(reason)))
            writer.writerow([key,patient,size,len(lines),duplicate_lines,conflicts,gaps,out_of_order,complete,identity_validated,';'.join(reason)])
            if export and complete and identity_validated:
                text='\n'.join(next(iter(lines[n])) for n in sorted(lines))
                export.write(json.dumps({'note_key':json.loads(key),'PAT_ID':patient,'note_date':day,
                    'anchor_status':flag,'note_type':note_type,'text':text},ensure_ascii=False)+'\n')
    db.commit()
    status = 'available' if verified else 'candidate_groups_only_identity_unverified'
    metrics_file(out,'notes_reconstruction_quality.csv',dict(quality),status)
    stats['reconstruction_quality'] = dict(quality)
    stats['unique_reconstructed_notes'] = quality['reconstructed_notes'] if verified else None
    stats['unique_validated_notes'] = quality['validated_note_groups'] if verified else None
    if not verified:
        report.issue('notes_identity','Candidate key granularity is not proof of note identity. NOTE_CSN may be encounter-level or version-level; same-day/same-type collisions can be undetectable without a verified note identifier/export contract.')
        report.skip('notes_reconstruction','Note identity unverified. Candidate-group quality audited; distinct-note counts and text reconstruction withheld. Configure locally verified note_key_fields and note_key_verified.')
    else: report.data['completed_modules'].append('notes_reconstruction')
    report.issue('notes_reconstruction','Contiguous LINE ordering is structural completeness only. Missing terminal lines cannot be detected without an export total/version contract. Exact duplicate fragments are collapsed for text; conflicting lines are never chosen.')
    return quality


def anchor_alignment(db, stats, reference, report, out):
    reliable = stats['patient_linkage_available'] and reference is not None
    verified = stats['note_identity_verified_by_local_configuration']
    status = 'available' if reliable else 'unavailable_patient_linkage_or_demographics'
    note_status = 'validated_notes' if verified else 'candidate_groups_identity_unverified'
    has_flag = 'anchor_flag' in stats['column_mapping']
    yes_rows = db.execute('SELECT COUNT(*) FROM fragments WHERE flag="Yes"').fetchone()[0]
    # Candidate dates still audit the row-level Yes flag when note identity is unverified.
    yes_dates = defaultdict(set)
    yes_patients = set()
    unknown_date_patients = set()
    for patient, day in db.execute('SELECT DISTINCT patient,day FROM fragments WHERE flag="Yes" AND patient IS NOT NULL'):
        yes_patients.add(patient)
        if day: yes_dates[patient].add(day)
        else: unknown_date_patients.add(patient)
    validated_patients = set()
    validated_dates = defaultdict(set)
    for patient,day in db.execute('SELECT DISTINCT patient,day FROM notes WHERE flag="Yes" AND validated=1 AND patient IS NOT NULL'):
        validated_patients.add(patient)
        if day: validated_dates[patient].add(day)
    candidates = Counter()
    candidate_days = defaultdict(set)
    for patient,day in db.execute('SELECT patient,day FROM notes WHERE flag="Yes" AND patient IS NOT NULL'):
        candidates[patient] += 1
        if day: candidate_days[patient].add(day)
    summary = {
        'anchor_yes_rows':yes_rows if has_flag else None,
        'distinct_anchor_yes_notes':db.execute('SELECT COUNT(*) FROM notes WHERE flag="Yes" AND validated=1').fetchone()[0] if verified and has_flag else None,
        'anchor_yes_candidate_groups':db.execute('SELECT COUNT(*) FROM notes WHERE flag="Yes"').fetchone()[0] if has_flag else None,
        'patients_with_anchor_yes_rows':len(yes_patients) if stats['patient_linkage_available'] and has_flag else None,
        'patients_with_anchor_yes_notes':db.execute('SELECT COUNT(DISTINCT patient) FROM notes WHERE flag="Yes" AND validated=1').fetchone()[0] if verified and has_flag and stats['patient_linkage_available'] else None,
        'patients_anchor_yes_dates_missing':len(unknown_date_patients) if stats['patient_linkage_available'] else None,
    }
    if not reliable:
        metrics_file(out,'notes_patient_intersections.csv',{'notes_source_patients':stats['unique_patients'], 'demographic_patients':len(reference['patients']) if reference else None},status)
        metrics_file(out,'notes_anchor_alignment_summary.csv',summary,status)
        metrics_file(out,'notes_multiple_anchor_summary.csv',{'demographic_multiple_calendar_anchor_patients':len(reference['ambiguous']) if reference else None},status)
        report.skip('notes_patient_intersections',status)
        report.skip('notes_anchor_alignment',status)
        return
    source = {row[0] for row in db.execute('SELECT DISTINCT patient FROM fragments WHERE patient IS NOT NULL')}
    demographic = reference['patients']
    intersections = {
        'notes_source_patients':len(source),'demographic_reference_patients':len(demographic),
        'notes_patients_matched_to_demographics':len(source & demographic),
        'notes_patients_outside_demographics':len(source-demographic),
        'demographics_patients_without_notes':len(demographic-source),
        'demographics_patients_with_anchor_yes_rows':len(demographic & yes_patients),
        'demographics_patients_with_anchor_yes_notes':len(demographic & validated_patients) if verified else None,
    }
    anchor_fields_available = all(field in stats['column_mapping'] for field in ['anchor_flag','date'])
    if 'anchor_flag' not in stats['column_mapping']:
        intersections['demographics_patients_with_anchor_yes_rows'] = None
        intersections['demographics_patients_with_anchor_yes_notes'] = None
    metrics_file(out,'notes_patient_intersections.csv',intersections)
    stats['patient_intersections']=intersections
    if not anchor_fields_available:
        metrics_file(out,'notes_anchor_alignment_summary.csv',{
            'mapped_anchor_flag':int('anchor_flag' in stats['column_mapping']),
            'mapped_note_date':int('date' in stats['column_mapping'])},'unavailable_anchor_flag_or_note_date')
        metrics_file(out,'notes_multiple_anchor_summary.csv',{
            'demographic_multiple_calendar_anchor_patients':len(reference['ambiguous'])},'unavailable_anchor_flag_or_note_date')
        report.data['completed_modules'].append('notes_patient_intersections')
        report.skip('notes_anchor_alignment','Anchor flag or note date missing/ambiguous; patient coverage available, alignment withheld.')
        return
    all_counts, multiple_counts = Counter(), Counter()
    validated_counts, validated_multiple_counts = Counter(), Counter()
    with (out/'notes_anchor_alignment_patient_audit.csv').open('w',newline='',encoding='utf-8') as handle:
        writer=csv.writer(handle)
        writer.writerow(['PAT_ID','demographic_anchor_days','anchor_yes_row_days','anchor_yes_candidate_note_days',
            'matching_days','row_date_sets_equal','anchor_yes_candidate_notes','matching_candidate_notes',
            'validated_matching_notes','validated_anchor_yes_note_dates','validated_matching_anchor_days','has_undated_anchor_yes_rows','multiple_demographic_anchor_days'])
        for patient in sorted(demographic):
            demographic_days=reference['dates'].get(patient,set())
            days=yes_dates.get(patient,set())
            matches=days & demographic_days
            match_notes=db.execute('''SELECT COUNT(*) FROM notes n JOIN reference r
                ON n.patient=r.patient AND n.day=r.day WHERE n.patient=? AND n.flag="Yes"''',(patient,)).fetchone()[0]
            validated_matches=db.execute('''SELECT COUNT(*) FROM notes n JOIN reference r
                ON n.patient=r.patient AND n.day=r.day WHERE n.patient=? AND n.flag="Yes" AND n.validated=1''',(patient,)).fetchone()[0] if verified else None
            has_invalid=patient in reference['invalid_anchor_patients'] or patient in unknown_date_patients
            values={
                'patients':1,'patients_with_anchor_yes_rows':int(patient in yes_patients),
                'patients_with_any_matching_anchor_day':int(bool(matches)),
                'patients_no_matching_anchor_day':int(not matches),
                'patients_exact_anchor_date_set_agreement':int(bool(days) and days==demographic_days and not has_invalid),
                'patients_partial_anchor_date_overlap':int(bool(matches) and days!=demographic_days),
                'patients_anchor_yes_dates_outside_demographic_dates':int(bool(days-demographic_days)),
                'patients_multiple_distinct_anchor_yes_dates':int(len(days)>1),
                'patients_multiple_candidate_anchor_notes':int(candidates[patient]>1),
                'patients_no_matching_candidate_anchor_note':int(match_notes==0),
                'patients_multiple_matching_candidate_anchor_notes':int(match_notes>1),
                'matching_anchor_days':len(matches),'unmatched_note_anchor_days':len(days-demographic_days),
                'unmatched_demographic_anchor_days':len(demographic_days-days),
                'patients_with_invalid_or_missing_anchor_dates':int(has_invalid),
            }
            valid_days=validated_dates.get(patient,set())
            valid_matches=valid_days & demographic_days
            valid_values={
                'patients_with_validated_anchor_yes_notes':int(patient in validated_patients),
                'patients_with_multiple_validated_anchor_note_dates':int(len(valid_days)>1),
                'patients_with_matching_validated_anchor_note':int(bool(validated_matches)),
                'patients_no_matching_validated_anchor_note':int(not validated_matches),
                'patients_multiple_matching_validated_anchor_notes':int(bool(validated_matches and validated_matches>1)),
                'patients_validated_anchor_date_set_agreement':int(bool(valid_days) and valid_days==demographic_days and not has_invalid),
                'validated_matching_anchor_days':len(valid_matches),
            }
            validated_counts.update(valid_values)
            if patient in reference['ambiguous']: validated_multiple_counts.update(valid_values)
            all_counts.update(values)
            if patient in reference['ambiguous']: multiple_counts.update(values)
            writer.writerow([patient,len(demographic_days),len(days),len(candidate_days.get(patient,set())),
                len(matches),bool(days) and days==demographic_days and not has_invalid,candidates[patient],match_notes,
                validated_matches,len(valid_days) if verified else None,len(valid_matches) if verified else None,patient in unknown_date_patients,patient in reference['ambiguous']])
    summary.update(all_counts)
    for metric in ['patients','patients_with_anchor_yes_rows','patients_with_any_matching_anchor_day',
        'patients_no_matching_anchor_day','patients_exact_anchor_date_set_agreement',
        'patients_partial_anchor_date_overlap','patients_anchor_yes_dates_outside_demographic_dates',
        'patients_multiple_distinct_anchor_yes_dates','patients_multiple_candidate_anchor_notes',
        'patients_no_matching_candidate_anchor_note','patients_multiple_matching_candidate_anchor_notes',
        'matching_anchor_days','unmatched_note_anchor_days','unmatched_demographic_anchor_days',
        'patients_with_invalid_or_missing_anchor_dates']:
        summary.setdefault(metric,0)
        multiple_counts.setdefault(metric,0)
    validated_metrics=['patients_with_validated_anchor_yes_notes','patients_with_multiple_validated_anchor_note_dates',
        'patients_with_matching_validated_anchor_note','patients_no_matching_validated_anchor_note',
        'patients_multiple_matching_validated_anchor_notes','patients_validated_anchor_date_set_agreement',
        'validated_matching_anchor_days']
    for metric in validated_metrics:
        summary[metric]=validated_counts[metric] if verified else None
        multiple_counts[metric]=validated_multiple_counts[metric] if verified else None
    summary['note_identity_status']=note_status
    multiple_counts['demographic_multiple_calendar_anchor_patients']=len(reference['ambiguous'])
    metrics_file(out,'notes_anchor_alignment_summary.csv',summary)
    metrics_file(out,'notes_multiple_anchor_summary.csv',dict(multiple_counts),note_status)
    stats['patient_intersections']=intersections
    stats['anchor_alignment']=dict(summary)
    stats['multiple_anchor_alignment']=dict(multiple_counts)
    report.data['completed_modules'].extend(['notes_patient_intersections','notes_anchor_alignment'])
    report.issue('notes_anchor_alignment','Anchor=Yes clinical meaning unconfirmed. Comparisons use calendar dates and all original Demographics anchors; no index visit, cohort membership or frailty label changed. Multiple Anchor=Yes notes can exist on one day.')
    plot(pd.Series({k:v for k,v in all_counts.items() if k in {'patients_with_any_matching_anchor_day','patients_no_matching_anchor_day','patients_multiple_candidate_anchor_notes'}},dtype=float),
         out/'figures'/'notes_anchor_alignment.png','Notes Anchor=Yes alignment (audit only)','Alignment',report=report)


def profile_outputs(db, stats, verified, types, flags, quality, report, out):
    status='available' if verified else 'note_identity_unverified'
    inventory={
        'total_rows':stats['row_count'],'unique_patients':stats['unique_patients'],
        'exact_duplicate_rows':stats['duplicate_rows'],
        'unique_validated_notes':stats['unique_validated_notes'],
        'unique_reconstructed_notes':stats['unique_reconstructed_notes'],
        'candidate_note_groups':quality['candidate_note_groups'],
        'missing_note_id_rows':stats['mapped_fields_missing_rows']['note_id'],
        'missing_note_csn_rows':stats['mapped_fields_missing_rows']['note_csn'],
        'missing_candidate_key_rows':quality['rows_without_candidate_key'],
        'missing_line_rows':quality['rows_missing_line'],'invalid_line_rows':quality['rows_invalid_line'],
        'missing_date_rows':quality['rows_missing_date'],'invalid_date_rows':quality['rows_invalid_date'],
        'missing_text_rows':quality['rows_missing_text'],'missing_anchor_rows':quality['rows_missing_anchor_indicator'],
        'unlinked_rows':stats['missing_patient_ids'],
        **stats['note_csn_granularity'],
    }
    note_metrics={'unique_validated_notes','unique_reconstructed_notes'}
    write_csv(out/'notes_inventory.csv',DELIVERABLES['notes_inventory.csv'],
        ((metric,value,status if metric in note_metrics else ('unavailable_linkage' if metric=='unique_patients' and value is None else 'available')) for metric,value in inventory.items()))
    counts=dict(db.execute('SELECT COALESCE(note_type,"(Missing)"),COUNT(*) FROM notes WHERE validated=1 GROUP BY note_type'))
    write_csv(out/'notes_type_distribution.csv',DELIVERABLES['notes_type_distribution.csv'],
              ((name,count,counts.get(name,0) if verified else None) for name,count in types.items()))
    note_flags=dict(db.execute('SELECT flag,COUNT(*) FROM notes WHERE validated=1 GROUP BY flag'))
    candidate_flags=dict(db.execute('SELECT flag,COUNT(*) FROM notes GROUP BY flag'))
    flag_status = 'available' if 'anchor_flag' in stats['column_mapping'] else 'unavailable_anchor_field'
    write_csv(out/'notes_anchor_summary.csv',DELIVERABLES['notes_anchor_summary.csv'],
        [(level,flag,count,label) for level,values,label in [('row',flags,flag_status),
            ('candidate_group',candidate_flags,'identity_not_implied' if flag_status=='available' else flag_status),('note',note_flags,status if flag_status=='available' else flag_status)]
         for flag,count in ((flag,values.get(flag,0) if level!='note' or (verified and flag_status=='available') else None) for flag in ['Yes','No','Missing','Other','Conflicting','Unresolved'])])
    db.execute('CREATE INDEX notes_patient ON notes(patient,flag,day)')
    with (out/'notes_per_patient_patient_audit.csv').open('w',newline='',encoding='utf-8') as handle:
        writer=csv.writer(handle)
        writer.writerow(['PAT_ID','fragment_rows','candidate_groups','validated_notes','reconstructed_notes'])
        for patient,rows in db.execute('SELECT patient,COUNT(*) FROM fragments WHERE patient IS NOT NULL GROUP BY patient'):
            candidates,notes,reconstructed=db.execute('''SELECT COUNT(*),COALESCE(SUM(validated),0),
                COALESCE(SUM(validated*structurally_complete),0) FROM notes WHERE patient=?''',(patient,)).fetchone()
            writer.writerow([patient,rows,candidates,notes if verified else None,reconstructed if verified else None])
    # Aggregate count distribution only, no patient IDs in figures.
    distribution=dict(db.execute('''SELECT count,COUNT(*) FROM
        (SELECT COUNT(*) AS count FROM notes WHERE validated=1 AND patient IS NOT NULL GROUP BY patient)
        GROUP BY count ORDER BY count''')) if verified else {}
    write_csv(out/'notes_per_patient_distribution.csv',['validated_notes_per_patient','patients'],distribution.items())
    if verified:
        plot(pd.Series(distribution,dtype=float),out/'figures'/'notes_per_patient.png','Validated Notes per linked patient','Notes',report=report)
    else:
        report.skip('figure_notes_per_patient','Note identity unverified; validated notes-per-patient distribution withheld.')
    plot(pd.Series(types,dtype=float).sort_values(ascending=False).head(25),out/'figures'/'notes_type_frequency.png','Notes types (fragment rows)','Note type',report=report)
    plot(pd.Series(flags,dtype=float),out/'figures'/'notes_anchor_rows.png','Notes Anchor flags (fragment rows)','Anchor flag',report=report)
    plot(pd.Series({k:quality[k] for k in ['structurally_complete_groups','failed_groups','groups_conflicting_lines','groups_out_of_order_fragments']} if quality['candidate_note_groups'] else {},dtype=float),
         out/'figures'/'notes_reconstruction_quality.png','Notes candidate-group reconstruction quality','Candidate group status',report=report)
    report.data['completed_modules'].append('notes_inventory')


def integrate_outputs(stats, reference, out):
    """Append the eighth dataset; never replace structured rows with Notes results."""
    path=out/'dataset_inventory.csv'
    existing=pd.read_csv(path) if path.exists() and path.stat().st_size>1 else pd.DataFrame()
    row={k:stats.get(k) for k in ['row_count','unique_patients','duplicate_rows','missing_patient_ids','unique_reconstructed_notes']}
    row['dataset']='notes'
    existing=existing.loc[existing.dataset.ne('notes')] if 'dataset' in existing else existing
    temporary=path.with_suffix('.notes-merge.tmp')
    pd.concat([existing,pd.DataFrame([row])],ignore_index=True).to_csv(temporary,index=False)
    temporary.replace(path)
    intersections=stats.get('patient_intersections')
    if intersections:
        path=out/'patient_set_intersections.csv'
        existing=pd.read_csv(path) if path.exists() and path.stat().st_size>1 else pd.DataFrame()
        existing=existing.loc[existing.dataset.ne('notes')] if 'dataset' in existing else existing
        row={'dataset':'notes','source_patients':intersections['notes_source_patients'],
            'demographic_patients':intersections['demographic_reference_patients'],
            'intersection_patients':intersections['notes_patients_matched_to_demographics'],
            'outside_demographics_patients':intersections['notes_patients_outside_demographics'],
            'demographic_patients_without_records':intersections['demographics_patients_without_notes'],
            'outside_demographics_rows':stats.get('outside_demographics_rows'),
            'missing_patient_id_rows':stats['missing_patient_ids'],
            'demographic_patients_with_anchor_yes_notes':intersections['demographics_patients_with_anchor_yes_notes']}
        temporary=path.with_suffix('.notes-merge.tmp')
        pd.concat([existing,pd.DataFrame([row])],ignore_index=True).to_csv(temporary,index=False)
        temporary.replace(path)


def audit_notes(root, config, report, out, demo=None):
    """Run independently; catch Notes errors without invalidating structured outputs."""
    settings = config.get('notes',{})
    report.data['notes_audit_settings']={k:settings.get(k) for k in [
        'note_key_fields','note_key_verified','line_start','encounter_linkage_verified','export_reconstructed_text']}
    try:
        with TemporaryDirectory(prefix='.notes-stage-',dir=out) as temporary:
            # sqlite3's transaction context alone does not close files on Windows.
            with closing(sqlite3.connect(str(Path(temporary)/'notes.sqlite'))) as db:
                db.execute('PRAGMA cache_size=-16384')
                # SQLite sort/index temporary files must remain in the approved output tree.
                quoted_temp=str(Path(temporary).resolve()).replace("'","''")
                db.execute("PRAGMA temp_store_directory='"+quoted_temp+"'")
                db.execute('PRAGMA temp_store=FILE')
                reference=prepare_reference(db,demo)
                stats,verified,types,flags=ingest(db,root,config,report,out)
                establish_linkage(db,stats,settings,reference,report)
                csn_granularity(db,stats,out)
                quality=reconstruct(db,stats,verified,settings,report,out)
                profile_outputs(db,stats,verified,types,flags,quality,report,out)
                # Count outside-reference rows separately without discarding audit input.
                if reference:
                    stats['outside_demographics_rows']=db.execute('''SELECT COUNT(*) FROM fragments f
                        WHERE patient IS NOT NULL AND NOT EXISTS (SELECT 1 FROM reference_patients r WHERE r.patient=f.patient)''').fetchone()[0]
                anchor_alignment(db,stats,reference,report,out)

    except Exception as exc:
        reason=str(exc) if isinstance(exc,PipelineError) else f'{type(exc).__name__}: Notes audit failed; review local schema/configuration. No note content logged.'
        report.fail('notes_audit',reason)
    finally:
        if 'notes' in report.data['datasets']:
            try: integrate_outputs(report.data['datasets']['notes'],None,out)
            except Exception as exc:
                report.fail('notes_integration',f'{type(exc).__name__}: Notes summary integration failed; structured outputs retained.')
        # Missing outputs are header-only with an explicit report reason, never invented counts.
        for filename,columns in DELIVERABLES.items():
            if not (out/filename).exists(): write_csv(out/filename,columns,[])
        known=set(report.data['completed_modules'])
        known.update(x['module'] for x in report.data['failed_modules'])
        known.update(x['module'] for x in report.data['skipped_modules'])
        for module in MODULES:
            if module not in known: report.skip(module,'Notes stage unavailable; see notes_audit failure or linkage/identity diagnostics.')
        for name in FIGURES:
            if 'figure_'+name not in known: report.skip('figure_'+name,'Notes prerequisite unavailable; see Notes module status.')
